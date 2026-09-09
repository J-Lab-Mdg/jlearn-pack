# -*- coding: utf-8 -*-
"""Génère SVT-9e-Manuel-Complet.docx — les 60 séances de l'année (B1 à B5),
numérotation continue 1-60, révisions et examens de chaque bimestre.
Normes J-Learn : fiche I/II/III, table 6 colonnes, Times New Roman,
code couleur maison (C00000/1E7B34/1F4E79/C2185B), sommaire interactif.
"""
import os, re, sys, zipfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from data_b1a import SEANCES_A as B1A
from data_b1b import SEANCES_B as B1B, REVISION_B1, EXAMEN_B1
from data_b2a import SEANCES_A as B2A
from data_b2b import SEANCES_B as B2B, REVISION_B2, EXAMEN_B2
from data_b3a import SEANCES_A as B3A
from data_b3b import SEANCES_B as B3B, REVISION_B3, EXAMEN_B3
from data_b4a import SEANCES_A as B4A
from data_b4b import SEANCES_B as B4B, REVISION_B4, EXAMEN_B4
from data_b5a import SEANCES_A as B5A
from data_b5b import SEANCES_B as B5B, REVISION_B5, EXAMEN_B5

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = "/home/user/jlearn-pack"
OUT = os.path.join(ROOT, "svt-9e", "output", "SVT-9e-Manuel-Complet.docx")
FONT = "Times New Roman"
ROUGE = RGBColor(0xC0, 0x00, 0x00)
VERT = RGBColor(0x1E, 0x7B, 0x34)
BLEU = RGBColor(0x1F, 0x4E, 0x79)
BORDEAUX = RGBColor(0xC2, 0x18, 0x5B)
NOIR = RGBColor(0x00, 0x00, 0x00)
BLEU_LIEN = RGBColor(0x05, 0x63, 0xC1)

BIMESTRES = [
    (1, B1A + B1B, REVISION_B1, EXAMEN_B1,
     "le corps humain, la division du corps et les organes, les membres supérieurs et la peau"),
    (2, B2A + B2B, REVISION_B2, EXAMEN_B2,
     "les maladies de la peau, les dents, l'hygiène des dents et l'œil"),
    (3, B3A + B3B, REVISION_B3, EXAMEN_B3,
     "le fonctionnement de l'œil, la protection et l'hygiène des yeux, l'oreille et la langue"),
    (4, B4A + B4B, REVISION_B4, EXAMEN_B4,
     "la santé et la sécurité du membre supérieur, les membres inférieurs et l'appareil respiratoire"),
    (5, B5A + B5B, REVISION_B5, EXAMEN_B5,
     "l'appareil digestif, le paludisme, les vers intestinaux, les maladies infantiles, l'habitat, la poule, le zébu, l'arbre, l'air et l'eau"),
]

# Numérotation continue 1..60
N = 0
for _nb, seances, _rev, _ex, _desc in BIMESTRES:
    for s in seances:
        N += 1
        s["n"] = N
TOTAL = N  # 60

_bm_id = [1]

# ---------------- helpers ----------------
def rich_runs(p, text, size=11, color=NOIR, bold=False, italic=False, kw_color=None, kw_bold=True):
    for chunk in re.split(r"(\*\*.+?\*\*)", text):
        if not chunk:
            continue
        is_kw = chunk.startswith("**") and chunk.endswith("**")
        r = p.add_run(chunk[2:-2] if is_kw else chunk)
        r.font.name = FONT
        r.font.size = Pt(size)
        if is_kw and kw_color is not None:
            r.font.color.rgb = kw_color
            r.font.bold = kw_bold
        else:
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.italic = italic
    return p

def para(doc, text, size=11, color=NOIR, bold=False, italic=False, align=None,
         kw_color=None, space_after=6, bookmark=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    rich_runs(p, text, size, color, bold, italic, kw_color)
    if bookmark:
        add_bookmark(p, bookmark)
    return p

def add_bookmark(p, name):
    i = _bm_id[0]; _bm_id[0] += 1
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(i)); s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(i))
    p._p.append(s); p._p.append(e)

def sommaire_link(doc, anchor, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    hl = OxmlElement("w:hyperlink"); hl.set(qn("w:anchor"), anchor); hl.set(qn("w:history"), "1")
    run = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts"); rf.set(qn("w:ascii"), FONT); rf.set(qn("w:hAnsi"), FONT); rPr.append(rf)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "22"); rPr.append(sz)
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rPr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    run.append(rPr)
    t = OxmlElement("w:t"); t.text = text; run.append(t)
    hl.append(run); p._p.append(hl)
    return p

def cell_add(cell, text, size=10, color=NOIR, bold=False, italic=False, kw_color=None, first=False):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    rich_runs(p, text, size, color, bold, italic, kw_color)
    return p

def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "none"); e.set(qn("w:sz"), "0")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), "auto")
        borders.append(e)
    tblPr.append(borders)

def set_widths(table, widths_cm):
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)

def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------------- document ----------------
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(2.0)
st = doc.styles["Normal"]
st.font.name = FONT; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

C = WD_ALIGN_PARAGRAPH.CENTER

# Couverture
para(doc, "Collection J-Learn", 16, NOIR, True, False, C, space_after=24)
para(doc, "Manuel de Sciences de la Vie et de la Terre", 22, ROUGE, True, False, C, space_after=12)
para(doc, "Classe de 9e — Année complète", 16, NOIR, True, False, C, space_after=12)
para(doc, "Programme officiel malgache — 60 séances, 5 bimestres", 12, NOIR, False, False, C, space_after=24)
para(doc, "Chaque séance : fiche de préparation + leçon illustrée + exercices corrigés.", 11, NOIR, False, True, C)
para(doc, "Manuel de fin d'année : il regroupe les cinq livraisons (B1 à B5) avec une numérotation continue des séances, de 1 à 60.", 11, NOIR, False, True, C)
page_break(doc)

# Avant-propos
para(doc, "Avant-propos", 16, NOIR, True, False, None, bookmark="bm_avantpropos")
para(doc, "Ce manuel de Sciences de la Vie et de la Terre pour la classe de 9e suit le programme officiel malgache. Il couvre toute l'année scolaire en cinq bimestres : le corps humain et la peau, les dents et l'œil, le fonctionnement de l'œil avec l'oreille et la langue, les membres et la respiration, puis la digestion, les grandes maladies, l'habitat et l'environnement.")
para(doc, "La méthode J-Learn organise chaque séance de 30 minutes en trois grandes étapes : la Révision, qui réactive les acquis de la séance précédente ; la Nouvelle Leçon, qui fait découvrir la notion du jour en six sous-étapes ; et l'Évaluation, qui vérifie ce qui est compris.")
para(doc, "Chaque séance comprend une fiche de préparation pour l'enseignant, une page de leçon illustrée et des exercices avec leur corrigé détaillé. À la fin de chaque bimestre, une révision et un sujet d'examen permettent de mesurer les progrès ; le Bimestre 5 se termine par la révision générale de fin d'année.")
para(doc, "Pour bien utiliser ce manuel, lire le mode d'emploi qui suit.")

# Mode d'emploi
para(doc, "Mode d'emploi", 16, NOIR, True, False, None, bookmark="bm_mode")
para(doc, "Comment est construite chaque séance ?", 12, NOIR, True)
para(doc, "I. Révision (5 min) — questions courtes sur la séance précédente, avec réponse attendue (R.A.).")
para(doc, "II. Nouvelle leçon (20 min) — six sous-étapes sans durée individuelle : 1. Mise en situation, 2. Présentation, 3. Observation, 4. Analyse, 5. Synthèse, 6. Application.")
para(doc, "III. Évaluation (5 min) — exercices écrits, avec corrigé.")
para(doc, "Après la fiche de préparation viennent la page LEÇON (le contenu à retenir, mots clés en bleu) puis la section EXERCICES (les mêmes exercices, notés sur 20, avec corrigé détaillé en bordeaux).")
para(doc, "Code couleur utilisé dans ce manuel :", 12, NOIR, True)
para(doc, "Titre de la leçon : rouge. Sous-titres : vert. Mots clés : bleu et gras. Corrigé : bordeaux. Texte courant : noir. Police : Times New Roman.")

# Sommaire
para(doc, "Sommaire", 16, NOIR, True, False, None, bookmark="bm_sommaire")
sommaire_link(doc, "bm_avantpropos", "Avant-propos")
sommaire_link(doc, "bm_mode", "Mode d'emploi")
sommaire_link(doc, "bm_tableau", "Tableau de bord de l'année")
for nb, seances, rev, exa, desc in BIMESTRES:
    para(doc, f"Bimestre {nb}", 13, VERT, True, space_after=2)
    sommaire_link(doc, f"bm_bimestre{nb}", f"Ouverture du Bimestre {nb}")
    for s in seances:
        sommaire_link(doc, f"bm_seance{s['n']}", f"Séance {s['n']} : {s['titre']}")
    sommaire_link(doc, f"bm_revision{nb}", rev["titre"])
    sommaire_link(doc, f"bm_examen{nb}", exa["titre"])

# Tableau de bord de l'année
page_break(doc)
para(doc, "Tableau de bord de l'année", 16, NOIR, True, False, None, bookmark="bm_tableau")
rows = []
for nb, seances, rev, exa, _desc in BIMESTRES:
    # Regrouper les thèmes consécutifs
    groups = []
    for s in seances:
        if groups and groups[-1][0] == s["theme"]:
            groups[-1][2] = s["n"]
        else:
            groups.append([s["theme"], s["n"], s["n"]])
    for theme, a, b in groups:
        rows.append((f"B{nb} — {theme}", f"Séances {a} à {b}" if a != b else f"Séance {a}"))
    rows.append((rev["titre"] + f" (B{nb})", f"Après la séance {seances[-1]['n']}"))
    rows.append((exa["titre"] + f" (B{nb})", "Après la révision"))
tb = doc.add_table(rows=1 + len(rows), cols=2, style="Table Grid")
set_widths(tb, [11.0, 6.0])
cell_add(tb.rows[0].cells[0], "Thème", bold=True, first=True)
cell_add(tb.rows[0].cells[1], "Séances", bold=True, first=True)
for i, (a, b) in enumerate(rows, start=1):
    cell_add(tb.rows[i].cells[0], a, first=True)
    cell_add(tb.rows[i].cells[1], b, first=True)

# ---------------- séances ----------------
def qa_paras_ens(cell, qa, first):
    cell_add(cell, "Questions :", bold=True, first=first)
    for i, (q, _a) in enumerate(qa, start=1):
        cell_add(cell, f"{i}. {q}")

def qa_paras_app(cell, qa, first):
    cell_add(cell, "R.A. :", bold=True, first=first)
    for i, (_q, a) in enumerate(qa, start=1):
        cell_add(cell, f"{i}. {a}")

def exo_paras_ens(cell, exos, first):
    f = first
    for e in exos:
        cell_add(cell, e["c"], bold=True, first=f); f = False
        for it in e["items"]:
            cell_add(cell, it)

def exo_paras_app(cell, exos, first):
    cell_add(cell, "Corrigé :", bold=True, first=first)
    for e in exos:
        for c in e["cor"]:
            cell_add(cell, c, kw_color=BORDEAUX)

for nb, seances, rev, exa, desc in BIMESTRES:
    # Page d'ouverture du bimestre
    page_break(doc)
    para(doc, f"BIMESTRE {nb}", 20, ROUGE, True, False, C, bookmark=f"bm_bimestre{nb}")
    para(doc, f"Ce bimestre couvre : {desc}.", 12, NOIR, False, True, C)
    para(doc, f"{len(seances)} séances, de la séance {seances[0]['n']} à la séance {seances[-1]['n']}.", 12, NOIR, False, False, C)

    for s in seances:
        n = s["n"]
        page_break(doc)
        para(doc, f"SÉANCE {n} / {TOTAL}", 14, NOIR, True, False, C, bookmark=f"bm_seance{n}")
        para(doc, s["titre"], 16, ROUGE, True, False, C)
        # Fiche
        para(doc, "FICHE DE PRÉPARATION", 14, NOIR, True, False, C)
        meta = doc.add_table(rows=6, cols=2)
        no_borders(meta)
        set_widths(meta, [11.5, 5.5])
        gauche = [f"Discipline : Sciences de la Vie et de la Terre", f"Thème : {s['theme']}",
                  f"Titre : {s['titre']}", f"Objectif spécifique : Être capable de {s['objectif']}.",
                  f"Documentation : {s['doc']}", "Support et matériel : voir le déroulement."]
        droite = ["Date : ...", "Classe : 9e", f"Séance n° : {n} / {TOTAL}", "Durée : 30 min", "", ""]
        for i in range(6):
            cell_add(meta.rows[i].cells[0], gauche[i], size=10, first=True)
            if droite[i]:
                cell_add(meta.rows[i].cells[1], droite[i], size=10, first=True)
        # Déroulement
        der = doc.add_table(rows=11, cols=6, style="Table Grid")
        set_widths(der, [2.4, 4.4, 4.4, 2.0, 2.0, 1.6])
        h0, h1 = der.rows[0], der.rows[1]
        h0.cells[0].merge(h1.cells[0]); h0.cells[1].merge(h0.cells[2])
        h0.cells[3].merge(h1.cells[3]); h0.cells[4].merge(h1.cells[4]); h0.cells[5].merge(h1.cells[5])
        cell_add(h0.cells[0], "Étapes et Durée", bold=True, first=True)
        cell_add(h0.cells[1], "Déroulement de la leçon", bold=True, first=True)
        cell_add(h0.cells[3], "Technique et Stratégie", bold=True, first=True)
        cell_add(h0.cells[4], "Support et Matériel", bold=True, first=True)
        cell_add(h0.cells[5], "Observation", bold=True, first=True)
        cell_add(h1.cells[1], "Enseignant", bold=True, first=True)
        cell_add(h1.cells[2], "Apprenants", bold=True, first=True)

        def row(r, etape, ens_fn, app_fn, tech, supp):
            c = der.rows[r].cells
            cell_add(c[0], etape, bold=True, first=True)
            ens_fn(c[1], True); app_fn(c[2], True)
            cell_add(c[3], tech, first=True); cell_add(c[4], supp, first=True)

        row(2, "I. Révision (5 min)", lambda cl, f: qa_paras_ens(cl, s["revision"], f),
            lambda cl, f: qa_paras_app(cl, s["revision"], f), "Questions-réponses", "—")
        sec2 = der.rows[3].cells
        sec2[0].merge(sec2[5])
        cell_add(sec2[0], "II. NOUVELLE LEÇON (20 min)", bold=True, first=True)
        row(4, "1. Mise en situation",
            lambda cl, f: (cell_add(cl, s["situation"], first=f), cell_add(cl, "Qui peut répondre à cette question ?")),
            lambda cl, f: cell_add(cl, "Écoutent et répondent.", first=f), "Travail collectif", "Tableau noir")
        row(5, "2. Présentation",
            lambda cl, f: [cell_add(cl, t, first=(f and i == 0)) for i, t in enumerate(s["presentation"])],
            lambda cl, f: cell_add(cl, "Écoutent.", first=f), "Exposé", "Tableau noir")
        row(6, "3. Observation",
            lambda cl, f: [cell_add(cl, t, first=(f and i == 0)) for i, t in enumerate(s["observation"])],
            lambda cl, f: cell_add(cl, "Observent, touchent et montrent.", first=f), "Observation dirigée", s["support_obs"])
        row(7, "4. Analyse", lambda cl, f: qa_paras_ens(cl, s["analyse"], f),
            lambda cl, f: qa_paras_app(cl, s["analyse"], f), "Questions-réponses", "Tableau noir")
        row(8, "5. Synthèse",
            lambda cl, f: (cell_add(cl, "Je lis la synthèse, vous écoutez :", bold=True, first=f), cell_add(cl, s["synthese"])),
            lambda cl, f: cell_add(cl, "Écoutent.", first=f), "Exposé", "Tableau noir")
        row(9, "6. Application", lambda cl, f: exo_paras_ens(cl, s["application"], f),
            lambda cl, f: exo_paras_app(cl, s["application"], f), "Travail individuel", "Cahier")
        row(10, "III. Évaluation (5 min)", lambda cl, f: exo_paras_ens(cl, s["evaluation"], f),
            lambda cl, f: exo_paras_app(cl, s["evaluation"], f), "Travail individuel", "Cahier")
        # Leçon
        page_break(doc)
        para(doc, "LEÇON", 14, NOIR, True, False, None)
        para(doc, s["titre"], 16, ROUGE, True, False, None, bookmark=f"bm_lecon{n}")
        if s["image"]:
            from PIL import Image as PILImage
            ipath = os.path.join(ROOT, s["image"][0])
            img = PILImage.open(ipath)
            iw, ih = img.size
            if max(iw, ih) > 800:
                r = 800.0 / max(iw, ih)
                img = img.convert("RGB").resize((int(iw * r), int(ih * r)), PILImage.LANCZOS)
            img = img.quantize(colors=256, method=PILImage.MEDIANCUT)
            optpath = os.path.join(os.path.dirname(OUT),
                                   os.path.splitext(os.path.basename(ipath))[0] + "_opt.png")
            img.save(optpath, optimize=True)
            pic = doc.add_paragraph(); pic.alignment = C
            pic.add_run().add_picture(optpath, width=Cm(10.0) if ih > iw else Cm(13.0))
            para(doc, s["image"][1], 10, NOIR, False, True, C)
        for sous_titre, partes in s["lecon"]:
            if sous_titre:
                para(doc, sous_titre, 12, VERT, True)
            for t in partes:
                para(doc, t, kw_color=BLEU)
        # Exercices
        para(doc, "EXERCICES", 14, NOIR, True)
        tous = s["application"] + s["evaluation"]
        for i, e in enumerate(tous, start=1):
            para(doc, f"Exercice {i} (5 points)", 11, NOIR, True)
            para(doc, e["c"], italic=False)
            for it in e["items"]:
                para(doc, it)
        para(doc, "Corrigé", 11, NOIR, True)
        for i, e in enumerate(tous, start=1):
            para(doc, f"Exercice {i}", 11, NOIR, True)
            for c in e["cor"]:
                para(doc, c, kw_color=BORDEAUX)
        para(doc, "Total : ... / 20", 11, NOIR, True)

    # Révision du bimestre
    page_break(doc)
    para(doc, rev["titre"], 16, ROUGE, True, False, None, bookmark=f"bm_revision{nb}")
    para(doc, "Je relis le tableau, puis je réponds aux questions.")
    tr = doc.add_table(rows=1 + len(rev["tableau"]), cols=2, style="Table Grid")
    set_widths(tr, [6.0, 11.0])
    cell_add(tr.rows[0].cells[0], "Notion", bold=True, first=True)
    cell_add(tr.rows[0].cells[1], "À retenir", bold=True, first=True)
    for i, (a, b) in enumerate(rev["tableau"], start=1):
        cell_add(tr.rows[i].cells[0], a, first=True)
        cell_add(tr.rows[i].cells[1], b, first=True)
    para(doc, "Questions de révision", 12, VERT, True)
    for i, (q, _a) in enumerate(rev["questions"], start=1):
        para(doc, f"{i}. {q}")
    para(doc, "Réponses attendues (R.A.)", 12, VERT, True)
    for i, (_q, a) in enumerate(rev["questions"], start=1):
        para(doc, f"{i}. {a}", kw_color=None)

    # Examen du bimestre
    page_break(doc)
    para(doc, exa["titre"], 16, ROUGE, True, False, None, bookmark=f"bm_examen{nb}")
    para(doc, f"Durée : {exa['duree']} — Total : ... / {exa['total']}", 11, NOIR, True)
    for i, e in enumerate(exa["parties"], start=1):
        para(doc, f"Partie {i}", 12, VERT, True)
        para(doc, e["c"])
        for it in e["items"]:
            para(doc, it)
    para(doc, "Corrigé du sujet d'examen", 14, NOIR, True)
    for i, e in enumerate(exa["parties"], start=1):
        para(doc, f"Partie {i}", 11, NOIR, True)
        for c in e["cor"]:
            para(doc, c, kw_color=BORDEAUX)
    para(doc, f"Total : ... / {exa['total']}", 11, NOIR, True)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("Sauvegardé :", OUT, os.path.getsize(OUT), "octets")

# ---------------- vérifications ----------------
z = zipfile.ZipFile(OUT)
xml = z.read("word/document.xml").decode("utf8")
tests = {
    "sectPr == 1": xml.count("<w:sectPr") == 1,
    "ns0 == 0": xml.count("ns0:") == 0,
    f"FICHES == {TOTAL}": xml.count("FICHE DE PRÉPARATION") == TOTAL,
    f"SÉANCE == {TOTAL}": len(re.findall(r"SÉANCE \d+ / %d" % TOTAL, xml)) == TOTAL,
    "Times New Roman > 0": xml.count("Times New Roman") > 0,
    ">Oral< == 0": xml.count(">Oral<") == 0,
    "BEPC == 0": xml.count("BEPC") == 0,
    "ministère == 0": xml.count("ministère") + xml.count("MINISTERE") == 0,
}
textes = re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml)
dbl = [t for t in textes if "  " in t]
tests["doubles espaces == 0"] = len(dbl) == 0
for k, v in tests.items():
    print(("OK  " if v else "ÉCHEC"), "-", k)
if dbl:
    print("Exemples :", dbl[:5])
bms = sorted(set(re.findall(r'w:name="(bm_[^\"]+)"', xml)))
print("Signets :", len(bms))
attendu = 4 + len(BIMESTRES) + 2 * TOTAL + 2 * len(BIMESTRES)
print("Signets attendus :", attendu)
