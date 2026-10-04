# Merge final — Manuel Mathématiques T8 (6 unités + annexes + sommaire complet)
from pathlib import Path
import re
from docx import Document
from docxcompose.composer import Composer
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

R = Path(__file__).resolve().parents[1]
W = R / 'work_merge'; W.mkdir(exist_ok=True)
(R / 'livrables').mkdir(exist_ok=True)
units = [R / f'Manuel_Mathematiques_T8_UNITE{i}.docx' for i in range(1, 7)]
ROMANS = {1: 'I', 2: 'II', 3: 'III', 4: 'IV', 5: 'V', 6: 'VI'}
NAMES = {1: 'Nombre', 2: 'Opération', 3: 'Algèbre', 4: 'Géométrie', 5: 'Mesure', 6: 'Traitement de données'}


def text(el):
    return ''.join(el.itertext()).strip()


def stripped(src, n):
    d = Document(src); body = d.element.body; found = False
    for el in list(body):
        if el.tag == qn('w:sectPr'):
            continue
        if el.tag == qn('w:p') and f'UNITÉ {ROMANS[n]} —' in text(el) and 'SÉANCE' not in text(el):
            found = True
        if not found:
            body.remove(el)
    if not found:
        raise RuntimeError(f'Unité {n} introuvable')
    out = W / f'u{n}.docx'; d.save(out); return out


def font(run, size=12, bold=False, color=None):
    run.font.name = 'Times New Roman'; run.font.size = Pt(size); run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def head(d, s, size=18):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(s), size, True, (192, 0, 0))


def para(d, s, bold=False):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    font(p.add_run(s), 12, bold); return p


# ---------- ANNEXES ----------
a = Document()
head(a, 'ANNEXES — MATHÉMATIQUES T8', 22)
para(a, 'Synthèse des notions étudiées dans les six unités.')
a.add_page_break(); head(a, 'ANNEXE 1 — GLOSSAIRE')
gloss = [
    ('Aire latérale', 'somme des aires des faces en pente d’un solide, sans la base'),
    ('Apothème', 'hauteur d’une face triangulaire d’une pyramide régulière'),
    ('Axe de symétrie', 'droite de pliage qui fait coïncider les deux moitiés d’une figure'),
    ('Caractère', 'question posée à chaque individu d’une population (couleur, taille, âge…)'),
    ('Centre de symétrie', 'point autour duquel un demi-tour ramène la figure sur elle-même'),
    ('Cône de révolution', 'solide engendré par un triangle rectangle qui tourne autour d’un côté de l’angle droit'),
    ('Effectif', 'nombre d’individus correspondant à une modalité'),
    ('Entier relatif', 'nombre entier muni d’un signe + ou −'),
    ('Équation', 'égalité contenant une inconnue à déterminer'),
    ('Événement contraire', 'tout ce qui n’est pas l’événement A ; P(Ā) = 1 − P(A)'),
    ('Extrapoler', 'estimer une valeur au-delà des données connues en prolongeant la tendance'),
    ('Fréquence', 'quotient de l’effectif d’une modalité par l’effectif total'),
    ('Génératrice', 'segment joignant le sommet d’un cône à un point du cercle de base'),
    ('Grandeur produit', 'grandeur obtenue en multipliant deux grandeurs : aire, énergie (kWh)'),
    ('Grandeur quotient', 'grandeur obtenue en divisant deux grandeurs : vitesse (km/h), prix au kilo (Ar/kg)'),
    ('Histogramme', 'diagramme à rectangles collés représentant des données groupées en classes'),
    ('Interpoler', 'estimer une valeur située entre deux données connues'),
    ('Modalité', 'réponse possible d’un caractère statistique'),
    ('Moyenne pondérée', 'moyenne tenant compte des effectifs ou des coefficients'),
    ('Notation scientifique', 'écriture a × 10ⁿ avec 1 ≤ a < 10'),
    ('Opposé', 'nombre symétrique par rapport à zéro : l’opposé de +4 est −4'),
    ('Patron', 'figure plane qui, pliée, reconstitue un solide'),
    ('Population', 'ensemble des individus concernés par une enquête statistique'),
    ('Probabilité', 'nombre entre 0 et 1 mesurant la chance qu’un événement se produise'),
    ('Proportion', 'égalité de deux rapports, vérifiable par les produits en croix'),
    ('Pyramide', 'solide à base polygonale dont les faces latérales triangulaires se rejoignent au sommet'),
    ('Rapport', 'quotient de deux quantités de même nature : nombre sans unité'),
    ('Symétrie centrale', 'demi-tour autour d’un point qui envoie M sur M′ tel que O soit le milieu de [MM′]'),
    ('Taux', 'quotient de deux grandeurs d’unités différentes : l’unité composée reste (Ar/kg, km/h)'),
    ('Volume', 'mesure de l’espace occupé par un solide, en unités cubes ; 1 L = 1 dm³'),
]
for t, v in gloss:
    p = a.add_paragraph()
    font(p.add_run(t + ' : '), 12, True, (31, 78, 121)); font(p.add_run(v + '.'), 12)

a.add_page_break(); head(a, 'ANNEXE 2 — FORMULES ET PROPRIÉTÉS')
rows = [
    ('Notion', 'Formule ou propriété'),
    ('Notation scientifique', 'a × 10ⁿ avec 1 ≤ a < 10 ; grand nombre → n positif, nombre < 1 → n négatif'),
    ('Règle des signes (× et ÷)', 'mêmes signes → + ; signes contraires → −'),
    ('Puissance d’un relatif', '(−a)ⁿ : positif si n pair, négatif si n impair'),
    ('Priorités', 'parenthèses → exposants → × ÷ → + − ; même priorité : de gauche à droite'),
    ('Produits en croix', 'a/b = c/d équivaut à a × d = b × c'),
    ('Proportionnalité', 'y = ax ; droite passant par l’origine'),
    ('Fonction affine', 'y = ax + b ; droite coupant l’axe vertical en b'),
    ('Symétries', 'conservent longueurs, angles, aires, alignement et parallélisme'),
    ('Aire latérale (pyramide régulière)', 'A = (périmètre de base × apothème) ÷ 2'),
    ('Aire latérale (cône)', 'A = π × r × g'),
    ('Volume de la pyramide', 'V = (aire de base × h) ÷ 3'),
    ('Volume du cône', 'V = (π × r² × h) ÷ 3'),
    ('Volumes et contenances', '1 L = 1 dm³ ; 1 mL = 1 cm³ ; 1 m³ = 1 000 L'),
    ('Vitesse', 'v = d ÷ t ; d = v × t ; t = d ÷ v'),
    ('Fréquence', 'effectif ÷ effectif total ; somme des fréquences = 1'),
    ('Angle d’un secteur circulaire', '(effectif ÷ total) × 360°'),
    ('Moyenne pondérée', 'somme des (valeur × effectif) ÷ effectif total'),
    ('Probabilité (équiprobabilité)', 'cas favorables ÷ cas possibles, entre 0 et 1'),
    ('Événement contraire', 'P(Ā) = 1 − P(A)'),
]
t = a.add_table(rows=0, cols=2); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, row in enumerate(rows):
    c = t.add_row().cells
    for j, s in enumerate(row):
        font(c[j].paragraphs[0].add_run(s), 11, i == 0)

a.add_page_break(); head(a, 'ANNEXE 3 — RÉSUMÉ DES MÉTHODES')
methods = [
    ('Écrire en notation scientifique', 'Placer la virgule après le premier chiffre non nul, compter les rangs de déplacement : recul → exposant positif, avance → exposant négatif.'),
    ('Comparer des relatifs', 'Positif > négatif toujours ; entre négatifs, le plus proche de zéro est le plus grand.'),
    ('Multiplier ou diviser des relatifs', 'Opérer sur les distances à zéro, puis appliquer la règle des signes ; vérifier par l’opération inverse.'),
    ('Vérifier une proportion', 'Calculer les deux produits en croix : s’ils sont égaux, les rapports forment une proportion ; sinon, non.'),
    ('Résoudre une équation', 'Isoler l’inconnue en faisant la même opération sur les deux membres — défaire le + b avant le × a — puis vérifier par substitution.'),
    ('Reconnaître la proportionnalité', 'Tableau : quotient y ÷ x constant ; graphique : points alignés avec l’origine.'),
    ('Construire un symétrique', 'Axiale : perpendiculaire à l’axe, même distance de l’autre côté. Centrale : viser le centre, traverser, même distance.'),
    ('Construire un patron', 'Dessiner la base exacte, attacher chaque face latérale à un côté de la base, contrôler les longueurs qui devront se recoller.'),
    ('Calculer un volume de pyramide ou de cône', 'Calculer l’aire de la base, multiplier par la hauteur, diviser par 3 ; convertir en litres si besoin (1 dm³ = 1 L).'),
    ('Comparer avec un taux unitaire', 'Ramener chaque offre à l’unité (1 kg, 1 h, 1 L) en divisant, puis comparer les taux obtenus.'),
    ('Traiter des données', 'Collecter, dresser le tableau des effectifs, calculer les fréquences, choisir le bon diagramme (bâtons, circulaire, histogramme).'),
    ('Calculer une probabilité', 'Vérifier l’équiprobabilité, compter les cas possibles puis les cas favorables, former le quotient et contrôler qu’il est entre 0 et 1.'),
]
for name, steps in methods:
    p = a.add_paragraph()
    font(p.add_run(name + ' — '), 12, True, (31, 78, 121)); font(p.add_run(steps), 12)

a.add_page_break(); head(a, 'ANNEXE 4 — AUTO-ÉVALUATION')
skills = [
    'Notation scientifique des grands et petits nombres',
    'Entiers relatifs : comparaison, opposé, addition',
    'Rapports, taux, proportions et produits en croix',
    'Multiplication, division et puissances de relatifs',
    'Priorités des opérations et calcul mental',
    'Opérations sur fractions et décimaux, estimation',
    'Proportionnalité y = ax et fonction affine y = ax + b',
    'Tableaux, graphiques, interpolation et extrapolation',
    'Équations du premier degré',
    'Symétries axiale et centrale : constructions et propriétés',
    'Programmes de construction et petites démonstrations',
    'Pyramide et cône : patrons, aires latérales, volumes',
    'Conversions, grandeurs produits et grandeurs quotients',
    'Statistiques : tableaux, diagrammes, fréquences, moyennes',
    'Probabilités et événement contraire',
]
t = a.add_table(rows=1, cols=4); t.style = 'Table Grid'
for i, s in enumerate(['Compétence', 'Acquis', 'En cours', 'À revoir']):
    font(t.rows[0].cells[i].paragraphs[0].add_run(s), 11, True)
for s in skills:
    c = t.add_row().cells
    font(c[0].paragraphs[0].add_run(s), 10)
    for i in range(1, 4):
        font(c[i].paragraphs[0].add_run('☐'), 14)

a.add_page_break(); head(a, 'ANNEXE 5 — INDEX')
para(a, ', '.join(sorted([x[0] for x in gloss], key=str.lower)) + '.')

a.add_page_break(); head(a, 'ANNEXE 6 — ÉVALUATIONS FORMAT EXAMEN')
para(a, 'Six sujets d’examen sur 20 points figurent à la fin des unités Nombre, Opération, Algèbre, Géométrie, Mesure et Traitement de données. Chaque sujet comporte cinq exercices avec corrigé détaillé et barème.')

a.add_page_break(); head(a, 'BIBLIOGRAPHIE ET WEBOGRAPHIE')
para(a, 'Ministère de l’Éducation Nationale de Madagascar, Programme d’études — Classe de T8, section Mathématiques (PE T8).')
para(a, 'Ministère de l’Éducation Nationale de Madagascar, Fascicule de ressources pédagogiques — Mathématiques.')
para(a, 'Collection J-Learn, Manuels de Mathématiques T6 et T7, J-Lab Madagascar, édition 2026.')
para(a, 'Collection J-Learn, charte de conception des manuels scolaires, version 18.')

annex_marks = {
    'ANNEXES — MATHÉMATIQUES T8': 'annexes', 'ANNEXE 1 — GLOSSAIRE': 'annexe1',
    'ANNEXE 2 — FORMULES ET PROPRIÉTÉS': 'annexe2', 'ANNEXE 3 — RÉSUMÉ DES MÉTHODES': 'annexe3',
    'ANNEXE 4 — AUTO-ÉVALUATION': 'annexe4', 'ANNEXE 5 — INDEX': 'annexe5',
    'ANNEXE 6 — ÉVALUATIONS FORMAT EXAMEN': 'annexe6', 'BIBLIOGRAPHIE ET WEBOGRAPHIE': 'sources',
}
mark_id = 5000
for paragraph in a.paragraphs:
    if paragraph.text in annex_marks:
        start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(mark_id)); start.set(qn('w:name'), annex_marks[paragraph.text])
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(mark_id))
        paragraph._p.insert(0, start); paragraph._p.append(end); mark_id += 1
ann = W / 'annexes.docx'; a.save(ann)

# ---------- COMPOSITION ----------
master = Document(units[0]); composer = Composer(master)
for i in range(2, 7):
    composer.append(Document(stripped(units[i - 1], i)))
composer.append(Document(ann))

# ---------- SOMMAIRE COMPLET ----------
def hyperlink_paragraph(label, anchor, indent=False):
    p = OxmlElement('w:p'); pPr = OxmlElement('w:pPr')
    spacing = OxmlElement('w:spacing'); spacing.set(qn('w:after'), '80'); pPr.append(spacing)
    if indent:
        ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '360'); pPr.append(ind)
    p.append(pPr)
    h = OxmlElement('w:hyperlink'); h.set(qn('w:anchor'), anchor); h.set(qn('w:history'), '1')
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '1F4E79')
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.extend([color, u]); r.append(rPr)
    tx = OxmlElement('w:t'); tx.text = label; r.append(tx); h.append(r); p.append(h)
    return p


def sentence(s, roman):
    if 'RÉVISION' in s:
        return f'Révision de l’unité {roman}'
    if 'EXAMEN' in s:
        return f'Sujet d’examen — unité {roman}'
    out = s.capitalize()
    return out


body_el = master.element.body
paras = list(master.paragraphs)
toc_i = next(i for i, p in enumerate(paras) if p.text == 'SOMMAIRE INTERACTIF')
insert_before = None
for par in paras[toc_i + 1:]:
    if par._p.xpath('.//w:br[@w:type="page"]'):
        insert_before = par._p; break

links = []
for unit_no in range(2, 7):
    links.append((f'Unité {ROMANS[unit_no]} — {NAMES[unit_no]}', f'u{unit_no}', False))
    source = Document(units[unit_no - 1])
    session_titles = [p.text for p in source.paragraphs if re.match(r'^SÉANCE \d+ / \d+ —', p.text)]
    for k, session_title in enumerate(session_titles, 1):
        raw = re.sub(r'^SÉANCE \d+ / \d+ — ', '', session_title)
        links.append((f'Séance {k} — {sentence(raw, ROMANS[unit_no])}', f'u{unit_no}_l{k}', True))
links.extend([
    ('Annexes', 'annexes', False), ('Annexe 1 — Glossaire', 'annexe1', True),
    ('Annexe 2 — Formules et propriétés', 'annexe2', True), ('Annexe 3 — Résumé des méthodes', 'annexe3', True),
    ('Annexe 4 — Auto-évaluation', 'annexe4', True), ('Annexe 5 — Index', 'annexe5', True),
    ('Annexe 6 — Évaluations format examen', 'annexe6', True), ('Bibliographie et webographie', 'sources', False),
])
if insert_before is not None:
    idx = body_el.index(insert_before)
    for label, anchor, ind in links:
        body_el.insert(idx, hyperlink_paragraph(label, anchor, ind)); idx += 1
else:
    raise RuntimeError('Point d’insertion du sommaire introuvable')

out = R / 'livrables' / 'Manuel_Mathematiques_T8_JLearn_V1.docx'
composer.save(out)
print('OK :', out, out.stat().st_size, 'octets')
