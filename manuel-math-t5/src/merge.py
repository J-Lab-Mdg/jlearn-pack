# Merge final — Manuel Mathématiques T5 (6 unités + annexes + sommaire complet)
from pathlib import Path
import re
from docx import Document
from docxcompose.composer import Composer
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

R = Path(__file__).resolve().parents[1]
W = R / 'work_merge'; W.mkdir(exist_ok=True)
(R / 'livrables').mkdir(exist_ok=True)
N_UNITS = 6
units = [R / f'Manuel_Mathematiques_T5_UNITE{i}.docx' for i in range(1, N_UNITS + 1)]
ROMANS = {1: 'I', 2: 'II', 3: 'III', 4: 'IV', 5: 'V', 6: 'VI'}
NAMES = {1: 'Nombre', 2: 'Opération', 3: 'Algèbre', 4: 'Géométrie', 5: 'Mesure', 6: 'Traitement de données'}


def text(el):
    return ''.join(el.itertext()).strip()


def stripped(src, n):
    d = Document(src); body = d.element.body; found = False
    for el in list(body):
        if el.tag == qn('w:sectPr'):
            continue
        if el.tag == qn('w:p') and f'UNITÉ {ROMANS[n]}' in text(el) and 'SÉANCE' not in text(el):
            found = True
        if not found:
            body.remove(el)
    if not found:
        raise RuntimeError(f'Unité {n} introuvable')
    out = W / f'u{n}.docx'; d.save(out); return out


def font(run, size=12, bold=False, color=None):
    run.font.name = 'Times New Roman'; run.font.size = Pt(size); run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def head(d, s, size=18):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(s), size, True, (192, 0, 0))


def para(d, s, bold=False):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    font(p.add_run(s), 12, bold); return p


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def card_para(cell, parts, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, fill=None, after=2):
    p = cell.add_paragraph(); p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    if fill:
        pPr = p._p.get_or_add_pPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); pPr.append(shd)
    for s, bold, color, italic in parts:
        r = p.add_run(s); font(r, size, bold, color); r.italic = italic
    return p


def annex_card(d, title, subtitle, image_name, formulas, units, example, note):
    table = d.add_table(rows=1, cols=1); table.style = 'Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cell, 'FFFFFF')
    card_para(cell, [(title, True, (31, 78, 121), False)], 15, WD_ALIGN_PARAGRAPH.CENTER, after=1)
    card_para(cell, [(subtitle, False, (102, 102, 102), False)], 9, WD_ALIGN_PARAGRAPH.CENTER, after=5)
    p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(6)
    run = p.add_run(); run.add_picture(str(R / 'assets' / 'figures' / image_name), width=Inches(5.9))
    for i, formula in enumerate(formulas if isinstance(formulas, list) else [formulas]):
        card_para(cell, [
            ('Formule : ' if i == 0 else '                 ', True, (31, 78, 121), False),
            (formula, True, (166, 28, 0), False)
        ], 11, WD_ALIGN_PARAGRAPH.LEFT, fill='EEF4F8', after=1)
    card_para(cell, [('Grandeurs et unité : ', True, (46, 125, 50), False), (units, False, (34, 34, 34), False)], 10, WD_ALIGN_PARAGRAPH.LEFT, fill='E8F5E9', after=2)
    card_para(cell, [('Exemple : ', True, (34, 34, 34), True), (example, False, (34, 34, 34), False)], 10, WD_ALIGN_PARAGRAPH.LEFT, after=2)
    card_para(cell, [(note, False, (34, 34, 34), False)], 9, WD_ALIGN_PARAGRAPH.LEFT, fill='FFF8E1', after=2)
    d.add_page_break()


# ---------- ANNEXES ----------
a = Document()
head(a, 'ANNEXES — MATHÉMATIQUES T5', 22)
para(a, 'Synthèse des notions étudiées dans les six unités, en préparation de l’examen du CEPE.')
a.add_page_break(); head(a, 'ANNEXE 1 — GLOSSAIRE')
gloss = [
    ('Aire', 'mesure, en unités carrées, de la surface d’une figure plane fermée'),
    ('Angle aigu', 'angle plus petit que l’angle droit'),
    ('Angle droit', 'angle formé par deux lignes perpendiculaires, comme le coin d’un cahier'),
    ('Angle obtus', 'angle plus grand que l’angle droit mais plus petit que l’angle plat'),
    ('Capacité', 'quantité de liquide qu’un récipient peut contenir ; unité principale : le litre'),
    ('Chiffre', 'signe qui sert à écrire les nombres : 0, 1, 2, 3, 4, 5, 6, 7, 8, 9'),
    ('Diagramme en barres', 'représentation des données par des barres dont la hauteur indique l’effectif'),
    ('Dividende', 'nombre que l’on partage dans une division'),
    ('Diviseur', 'nombre par lequel on divise'),
    ('Dixième', 'une des dix parts égales de l’unité ; premier chiffre après la virgule'),
    ('Centième', 'une des cent parts égales de l’unité ; deuxième chiffre après la virgule'),
    ('Effectif', 'nombre de fois qu’une donnée se présente dans une enquête'),
    ('Enquête', 'collecte de données auprès de personnes pour répondre à une question'),
    ('Équation', 'énoncé mathématique qui comporte une inconnue et le signe d’égalité'),
    ('Essais systématiques', 'méthode qui consiste à essayer des valeurs l’une après l’autre puis à vérifier'),
    ('Fraction', 'écriture qui représente une ou plusieurs parts égales de l’unité : numérateur sur dénominateur'),
    ('Fraction décimale', 'fraction dont le dénominateur est 10 ou 100'),
    ('Fraction impropre', 'fraction dont le numérateur est plus grand que le dénominateur'),
    ('Fractions équivalentes', 'fractions qui représentent la même quantité, comme 1/2 et 2/4'),
    ('Masse', 'ce que pèse un objet ; unité principale : le gramme et ses multiples'),
    ('Motif', 'ensemble des éléments qui se répètent dans une suite ; il peut être répété ou croissant'),
    ('Multiplicande', 'nombre que l’on multiplie'),
    ('Multiplicateur', 'nombre de fois que l’on prend le multiplicande'),
    ('Nombre décimal', 'nombre qui comporte une partie entière et une partie décimale séparées par une virgule'),
    ('Nombre fractionnaire', 'nombre formé d’un entier et d’une fraction, comme 1 et 1/2'),
    ('Périmètre', 'longueur du tour complet d’une figure plane'),
    ('Pictogramme', 'représentation des données par de petits dessins qui valent chacun une quantité'),
    ('Polygone', 'figure plane fermée dont les côtés sont des segments de droite'),
    ('Polygone régulier', 'polygone dont tous les côtés et tous les angles sont égaux, comme le carré'),
    ('Probabilité d’un événement', 'rapport du nombre de résultats favorables au nombre de résultats possibles'),
    ('Quadrilatère', 'polygone à quatre côtés'),
    ('Quotient', 'résultat de la division'),
    ('Reste', 'ce qui reste après une division quand le partage ne tombe pas juste'),
    ('Règle de trois', 'méthode pour trouver une quatrième valeur quand trois valeurs proportionnelles sont connues'),
    ('Suite', 'liste d’éléments arrangés dans un ordre déterminé'),
    ('Tableau des effectifs', 'tableau utilisé pour dénombrer les données et noter combien de fois chacune se présente'),
    ('Taux unitaire', 'valeur d’une seule unité, par exemple le prix d’un kilo de riz'),
    ('Terme', 'chacun des éléments d’une suite, d’une somme ou d’une différence'),
    ('Triangle équilatéral', 'triangle dont les trois côtés ont la même longueur'),
    ('Triangle isocèle', 'triangle qui a deux côtés de même longueur'),
    ('Triangle rectangle', 'triangle qui possède un angle droit'),
    ('Triangle scalène', 'triangle dont les trois côtés ont des longueurs toutes différentes'),
    ('Valeur de position', 'valeur d’un chiffre selon sa place dans le nombre : unités, dizaines, centaines…'),
    ('Volume', 'place qu’occupe un objet dans l’espace, mesurée en unités cubes (cm³)'),
]
for t, v in gloss:
    p = a.add_paragraph()
    font(p.add_run(t + ' : '), 12, True, (31, 78, 121)); font(p.add_run(v + '.'), 12)

sections = [
    ('PÉRIMÈTRES ET AIRES', 'On mesure le contour en unités de longueur et la surface en unités carrées. Les dimensions doivent être dans la même unité.', [
        ('Carré', 'Les quatre côtés ont la même longueur.', 'ann2_carre.png', ['P = 4 × c ;   A = c × c', 'c = P ÷ 4'], 'c en cm ou m ; P en cm ou m ; A en cm² ou m².', 'Si c = 5 cm : P = 20 cm ; A = 25 cm². Si P = 20 cm : c = 20 ÷ 4 = 5 cm.', 'Le côté c se retrouve avec le périmètre : c = P ÷ 4.'),
        ('Rectangle', 'Deux longueurs L et deux largeurs l.', 'ann2_rectangle.png', ['P = 2 × (L + l) ;   A = L × l', 'L = A ÷ l ;   l = A ÷ L', 'L = P ÷ 2 − l ;   l = P ÷ 2 − L'], 'L et l en cm ou m ; P en cm ou m ; A en cm² ou m².', 'Si L = 8 cm et l = 3 cm : P = 22 cm ; A = 24 cm². Si A = 24 cm² et l = 3 cm : L = 8 cm.', 'Les deux formules inverses du périmètre sont à connaître : L = P ÷ 2 − l et l = P ÷ 2 − L.'),
        ('Triangle', 'Trois côtés a, b et c ; une base b et sa hauteur h.', 'ann2_triangle.png', ['P = a + b + c ;   A = (b × h) ÷ 2'], 'P en cm ou m ; A en cm² ou m² ; h perpendiculaire à b.', 'Si a = 5 cm, b = 6 cm, c = 7 cm : P = 18 cm. Si h = 4 cm : A = 12 cm².', 'La somme des angles d’un triangle est 180°.'),
        ('Polygone régulier et demi-périmètre', 'Un polygone régulier possède n côtés égaux de longueur c.', 'ann2_polygone.png', ['P = n × c ;   DP = P ÷ 2'], 'P et DP en cm ou m ; n est sans unité.', 'Un hexagone de côté 4 cm : P = 6 × 4 = 24 cm ; DP = 12 cm.', 'Pour un polygone quelconque, on additionne tous les côtés.')
    ]),
    ('CERCLE ET FIGURES À BASES PARALLÈLES', 'Pour le cercle, on utilise π ≈ 3,14. Le diamètre traverse le centre ; le rayon va du centre au bord.', [
        ('Cercle et disque', 'Rayon r, diamètre d, circonférence C et aire du disque A.', 'ann2_cercle.png', ['d = 2r ; r = d ÷ 2 ; C = πd = 2πr ; A = πr²'], 'r et d en cm ou m ; C en cm ou m ; A en cm² ou m².', 'Si r = 3 cm : d = 6 cm ; C ≈ 18,84 cm ; A ≈ 28,26 cm².', 'Ne pas confondre la circonférence C, qui est une longueur, et l’aire A, qui est une surface.'),
        ('Trapèze', 'Grande base B, petite base b et hauteur h.', 'ann2_trapeze.png', ['A = (B + b) × h ÷ 2'], 'B, b et h en cm ou m ; A en cm² ou m².', 'Si B = 10 cm, b = 6 cm et h = 4 cm : A = 32 cm².', 'Formules inverses : h = (2A) ÷ (B + b) ; B = (2A ÷ h) − b.'),
        ('Parallélogramme', 'Base b et hauteur h perpendiculaire à cette base.', 'ann2_parallelogramme.png', ['A = b × h'], 'b et h en cm ou m ; A en cm² ou m².', 'Si b = 7 cm et h = 3 cm : A = 21 cm².', 'Formules inverses : b = A ÷ h ; h = A ÷ b.')
    ]),
    ('SOLIDES ET VOLUMES', 'Le volume mesure l’espace occupé. L’aire latérale couvre les faces autour du solide ; l’aire totale comprend aussi les bases.', [
        ('Cube', 'Les trois dimensions sont égales à l’arête a.', 'ann2_cube.png', ['A_latérale = 4 × a² ;   A_totale = 6 × a²', 'V = a × a × a = a³'], 'a en cm ou m ; aires en cm² ou m² ; V en cm³ ou m³.', 'Si a = 4 cm : A_latérale = 64 cm² ; A_totale = 96 cm² ; V = 64 cm³.', 'Un cube est un pavé droit particulier.'),
        ('Pavé droit', 'Longueur L, largeur l et hauteur h.', 'ann2_pave.png', ['A_latérale = 2 × (L + l) × h', 'A_totale = 2 × (L × l + L × h + l × h)', 'V = L × l × h'], 'L, l et h en cm ou m ; aires en cm² ou m² ; V en cm³ ou m³.', 'Si L = 8 cm, l = 3 cm et h = 2 cm : A_latérale = 44 cm² ; A_totale = 92 cm² ; V = 48 cm³.', 'La base est le rectangle L × l ; l’aire latérale entoure cette base.'),
        ('Prisme droit', 'Aire et périmètre de la base, puis hauteur du prisme.', 'ann2_prisme.png', ['A_latérale = P_base × h', 'A_totale = A_latérale + 2 × A_base', 'V = A_base × h'], 'P_base en cm ou m ; A_base, A_latérale et A_totale en cm² ou m² ; V en cm³ ou m³.', 'Si P_base = 14 cm, A_base = 12 cm² et h = 5 cm : A_latérale = 70 cm² ; A_totale = 94 cm² ; V = 60 cm³.', 'La base peut être un triangle, un rectangle ou un autre polygone.'),
        ('Cylindre', 'Rayon r de la base et hauteur h.', 'ann2_cylindre.png', ['A_latérale = 2 × π × r × h', 'A_totale = A_latérale + 2 × π × r²', 'V = π × r² × h'], 'r et h en cm ou m ; aires en cm² ou m² ; V en cm³ ou m³.', 'Si r = 2 cm et h = 5 cm : A_latérale ≈ 62,8 cm² ; A_totale ≈ 87,92 cm² ; V ≈ 62,8 cm³.', 'La base du cylindre est un disque : A_base = πr².')
    ]),
    ('PROPRIÉTÉS ET UNITÉS À RETENIR', 'Ces repères permettent de vérifier si une formule et une unité conviennent au problème.', [
        ('Angles des polygones étudiés', 'Les angles intérieurs se lisent sur la figure.', 'ann2_proprietes.png', ['Triangle : 180° ;   quadrilatère : 360°'], 'Un angle se mesure en degrés (°).', 'Dans un triangle, si deux angles mesurent 60° et 70°, le troisième mesure 50°.', 'Carré et rectangle : quatre angles droits.'),
        ('Choisir la bonne unité', 'La grandeur demandée indique l’unité du résultat.', 'ann2_unites.png', ['Longueur / périmètre : cm ou m ; aire : cm² ou m² ; volume : cm³ ou m³'], 'Capacité : mL, cL, dL, L ; 1 dm³ = 1 L ; 1 cm³ = 1 mL.', 'Une longueur de 8 cm ne peut pas donner une aire de 8 cm : il faut appliquer la formule et écrire cm².', 'Toujours convertir les mesures dans la même unité avant de calculer.')
    ])
]
a.add_page_break()
head(a, 'ANNEXE 2 — FORMULES ET PROPRIÉTÉS')
para(a, 'Chaque carte présente une figure repérée, les grandeurs utiles, la formule directement dessous, les unités et un exemple. Les lettres utilisées : c = côté, L = longueur, l = largeur, b = petite base, B = grande base, h = hauteur, r = rayon, d = diamètre, n = nombre de côtés.')
for section_title, intro, cards in sections:
    a.add_page_break()
    head(a, section_title, 16)
    para(a, intro)
    for title, subtitle, image_name, formulas, unit_desc, example, note in cards:
        annex_card(a, title, subtitle, image_name, formulas, unit_desc, example, note)

a.add_page_break(); head(a, 'ANNEXE 3 — RÉSUMÉ DES MÉTHODES')
methods = [
    ('Lire un grand nombre', 'Séparer les classes par tranches de trois chiffres en partant de la droite, puis lire classe des mille puis classe des unités.'),
    ('Comparer deux nombres', 'Compter les chiffres : le plus long gagne ; à longueur égale, comparer chiffre par chiffre depuis la gauche.'),
    ('Trouver une fraction équivalente', 'Multiplier ou diviser le haut et le bas par le même nombre, et vérifier avec les bandes de fractions.'),
    ('Poser une addition ou une soustraction', 'Aligner les unités sous les unités, les dizaines sous les dizaines ; ne pas oublier la retenue ; vérifier par l’opération inverse.'),
    ('Poser une multiplication', 'Multiplier par les unités, puis par les dizaines en décalant d’un rang, et additionner les produits partiels.'),
    ('Poser une division', 'Prendre les chiffres du dividende un à un, chercher combien de fois le diviseur entre, écrire le reste ; contrôler : diviseur × quotient + reste = dividende.'),
    ('Respecter la priorité des opérations', 'Calculer d’abord dans les parenthèses, ensuite les × et ÷ de gauche à droite, enfin les + et −.'),
    ('Calculer mentalement', 'Décomposer les nombres, utiliser les doubles et les moitiés, compenser (ajouter d’un côté, retirer de l’autre).'),
    ('Continuer une suite', 'Observer la régularité (ce qui se répète ou ce qui augmente), décrire le motif, puis prolonger et vérifier.'),
    ('Utiliser la règle de trois', 'Trouver d’abord ce que valent les quantités connues, multiplier puis diviser : (9 × 700) ÷ 3 pour 9 stylos.'),
    ('Trouver une valeur inconnue', 'Essayer des valeurs l’une après l’autre (essais systématiques), puis vérifier l’égalité en remplaçant l’inconnue.'),
    ('Reconnaître un triangle', 'Mesurer les côtés (équilatéral, isocèle, scalène) et observer les angles (rectangle, aigu, obtus).'),
    ('Reproduire une figure', 'Repérer les sommets sur le quadrillage, compter les carreaux, tracer à la règle et vérifier avec l’équerre.'),
    ('Convertir une unité', 'Placer le nombre dans le tableau de conversion, un chiffre par colonne, puis déplacer la virgule vers la nouvelle unité.'),
    ('Calculer un périmètre ou une aire', 'Choisir la bonne formule, remplacer par les mesures dans la même unité, calculer, écrire l’unité du résultat (cm ou cm²).'),
    ('Résoudre un problème', 'Lire deux fois l’énoncé, repérer la question et les données, choisir les opérations, calculer, puis vérifier que la réponse a du sens.'),
    ('Mener une enquête', 'Poser une bonne question, recueillir les réponses, les compter dans un tableau des effectifs, puis les représenter en diagramme.'),
    ('Estimer une probabilité', 'Répéter l’expérience plusieurs fois, compter les résultats dans un tableau, puis former le rapport des cas favorables au total.'),
]
for name, steps in methods:
    p = a.add_paragraph()
    font(p.add_run(name + ' — '), 12, True, (31, 78, 121)); font(p.add_run(steps), 12)

a.add_page_break(); head(a, 'ANNEXE 4 — AUTO-ÉVALUATION')
skills = [
    'Lire, écrire et comparer les nombres jusqu’à 1 000 000',
    'Décomposer et composer les grands nombres',
    'Fractions : représentation, équivalences, comparaison',
    'Fraction impropre et nombre fractionnaire',
    'Nombres décimaux jusqu’aux centièmes',
    'Fractions décimales et nombres décimaux',
    'Addition et soustraction jusqu’à 1 000 000',
    'Multiplication (3 chiffres × 2 chiffres)',
    'Division avec quotient et reste',
    'Multiplication et division de fractions',
    'Opérations sur les nombres décimaux',
    'Priorité des opérations',
    'Calcul mental : ×10, ×100, ×1 000, doubles, moitiés',
    'Suites, régularités, rang et terme',
    'Proportionnalité intuitive et règle de trois',
    'Égalités et valeurs inconnues',
    'Angles et classification des triangles',
    'Figures planes et outils de géométrie',
    'Composition et décomposition de polygones',
    'Longueur, masse, capacité : conversions',
    'Mesure du temps : conversions et opérations',
    'Périmètres, aires et volume intuitif',
    'Enquête, tableaux et diagrammes',
    'Probabilité expérimentale',
]
t = a.add_table(rows=1, cols=4); t.style = 'Table Grid'
for i, s in enumerate(['Compétence', 'Acquis', 'En cours', 'À revoir']):
    font(t.rows[0].cells[i].paragraphs[0].add_run(s), 11, True)
for s in skills:
    c = t.add_row().cells
    font(c[0].paragraphs[0].add_run(s), 10)
    for i in range(1, 4):
        font(c[i].paragraphs[0].add_run('☐'), 14)

a.add_page_break(); head(a, 'ANNEXE 5 — INDEX')
para(a, ', '.join(sorted([x[0] for x in gloss], key=str.lower)) + '.')

a.add_page_break(); head(a, 'ANNEXE 6 — ÉVALUATIONS FORMAT CEPE')
para(a, 'Six sujets d’examen sur 20 points figurent à la fin des unités Nombre, Opération, Algèbre, Géométrie, Mesure et Traitement de données. Chaque sujet comporte cinq exercices avec corrigé détaillé et barème, dans l’esprit des épreuves de calcul du CEPE : opérations posées, conversions et problème de la vie courante.')

a.add_page_break(); head(a, 'BIBLIOGRAPHIE ET WEBOGRAPHIE')
para(a, 'Ministère de l’Éducation Nationale de Madagascar, Programme d’études — Classe de T5, section Mathématiques (PE T5).')
para(a, 'Sujets d’examen du CEPE, épreuve de calcul, sessions antérieures, Madagascar.')
para(a, 'Collection J-Learn, Manuels de Mathématiques T6, T7, T8 et T9, J-Lab Madagascar, édition 2026.')
para(a, 'Collection J-Learn, charte de conception des manuels scolaires, version 18.')

annex_marks = {
    'ANNEXES — MATHÉMATIQUES T5': 'annexes', 'ANNEXE 1 — GLOSSAIRE': 'annexe1',
    'ANNEXE 2 — FORMULES ET PROPRIÉTÉS': 'annexe2', 'ANNEXE 3 — RÉSUMÉ DES MÉTHODES': 'annexe3',
    'ANNEXE 4 — AUTO-ÉVALUATION': 'annexe4', 'ANNEXE 5 — INDEX': 'annexe5',
    'ANNEXE 6 — ÉVALUATIONS FORMAT CEPE': 'annexe6', 'BIBLIOGRAPHIE ET WEBOGRAPHIE': 'sources',
}
mark_id = 5000
for paragraph in a.paragraphs:
    if paragraph.text in annex_marks:
        start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(mark_id)); start.set(qn('w:name'), annex_marks[paragraph.text])
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(mark_id))
        paragraph._p.insert(0, start); paragraph._p.append(end); mark_id += 1
ann = W / 'annexes.docx'; a.save(ann)

# ---------- COMPOSITION ----------
master = Document(units[0]); composer = Composer(master)
for i in range(2, N_UNITS + 1):
    composer.append(Document(stripped(units[i - 1], i)))
composer.append(Document(R / 'Manuel_Mathematiques_T5_METHODES.docx'))
composer.append(Document(R / 'Manuel_Mathematiques_T5_EXAMENS.docx'))
composer.append(Document(ann))

# ---------- SOMMAIRE COMPLET ----------
def hyperlink_paragraph(label, anchor, indent=False):
    p = OxmlElement('w:p'); pPr = OxmlElement('w:pPr')
    spacing = OxmlElement('w:spacing'); spacing.set(qn('w:after'), '80'); pPr.append(spacing)
    if indent:
        ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '360'); pPr.append(ind)
    p.append(pPr)
    h = OxmlElement('w:hyperlink'); h.set(qn('w:anchor'), anchor); h.set(qn('w:history'), '1')
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '1F4E79')
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.extend([color, u]); r.append(rPr)
    tx = OxmlElement('w:t'); tx.text = label; r.append(tx); h.append(r); p.append(h)
    return p


def sentence(s, roman):
    if 'RÉVISION' in s:
        return f'Révision de l’unité {roman}'
    if 'EXAMEN' in s:
        return f'Sujet d’examen — unité {roman}'
    return s.capitalize()


body_el = master.element.body
paras = list(master.paragraphs)
toc_i = next(i for i, p in enumerate(paras) if p.text == 'SOMMAIRE INTERACTIF')
insert_before = None
for par in paras[toc_i + 1:]:
    if par._p.xpath('.//w:br[@w:type="page"]'):
        insert_before = par._p; break

links = []
for unit_no in range(2, N_UNITS + 1):
    links.append((f'Unité {ROMANS[unit_no]} — {NAMES[unit_no]}', f'u{unit_no}', False))
    source = Document(units[unit_no - 1])
    session_titles = [p.text for p in source.paragraphs if re.match(r'^SÉANCE \d+ / \d+ —', p.text)]
    for k, session_title in enumerate(session_titles, 1):
        raw = re.sub(r'^SÉANCE \d+ / \d+ — ', '', session_title)
        links.append((f'Séance {k} — {sentence(raw, ROMANS[unit_no])}', f'u{unit_no}_l{k}', True))
links.extend([
    ('Méthodes d’or — La boîte à outils du CEPE (9 fiches)', 'mor', False),
    ('Sujets d’examen — Opérations (20 sujets corrigés)', 'exop', False),
    ('Sujets d’examen — Problèmes (20 sujets corrigés)', 'expb', False),
    ('Annexes', 'annexes', False), ('Annexe 1 — Glossaire', 'annexe1', True),
    ('Annexe 2 — Formules et propriétés', 'annexe2', True), ('Annexe 3 — Résumé des méthodes', 'annexe3', True),
    ('Annexe 4 — Auto-évaluation', 'annexe4', True), ('Annexe 5 — Index', 'annexe5', True),
    ('Annexe 6 — Évaluations format CEPE', 'annexe6', True), ('Bibliographie et webographie', 'sources', False),
])
if insert_before is not None:
    idx = body_el.index(insert_before)
    for label, anchor, ind in links:
        body_el.insert(idx, hyperlink_paragraph(label, anchor, ind)); idx += 1
else:
    raise RuntimeError('Point d’insertion du sommaire introuvable')

out = R / 'livrables' / 'Manuel_Mathematiques_T5_JLearn_V1.docx'
composer.save(out)

# ---------- allègement : recompression des JPEG embarqués ----------
import zipfile, shutil, io
from PIL import Image
tmp = R / 'work_merge' / 'light.docx'
with zipfile.ZipFile(out) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename.startswith('word/media/') and item.filename.lower().endswith(('.jpg', '.jpeg')) and len(data) > 80_000:
            im = Image.open(io.BytesIO(data)).convert('RGB')
            if max(im.size) > 1150:
                im.thumbnail((1150, 1150), Image.LANCZOS)
            buf = io.BytesIO(); im.save(buf, 'JPEG', quality=66, optimize=True)
            if buf.tell() < len(data):
                data = buf.getvalue()
        zout.writestr(item, data)
shutil.move(tmp, out)
print('OK :', out, out.stat().st_size, 'octets')
