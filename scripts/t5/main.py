# -*- coding: utf-8 -*-
"""Assemble 'ST T5 [PE] Fiche de preparation sujet corriges J-Learn.docx'.
Builds: cover + avant-propos + mode d'emploi + full global TOC + 11 unites
(seances 1-61) + ANNEXES, from content_unit1..content_unit11.py.
"""
import os
import sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import docx
import builder as B
from plan import UNITS, flat_seances, TOTAL_SEANCES

UNIT_MODULES = []
for i in range(1, 12):
    mod = __import__(f"content_unit{i}")
    UNIT_MODULES.append(mod)

OUT = os.path.join(os.path.dirname(os.path.dirname(HERE)),
                    "ST T5 [PE] Fiche de preparation sujet corriges J-Learn.docx")

ALL_SEANCES = flat_seances()


def build_lecon(doc, s):
    B.lecon_title_repeat(doc, s["title"])
    if s.get("cover_image"):
        path, caption = s["cover_image"]
        if os.path.exists(path):
            B.add_image(doc, path, caption)
        else:
            print("WARNING: missing image, skipped:", path)
    for item in s["lecon"]:
        kind, text = item[0], item[1]
        if kind == "section":
            B.section_header(doc, text)
        elif kind == "sub":
            B.sub_header(doc, text)
        elif kind == "body":
            if isinstance(text, str):
                for para_text in text.split("\n"):
                    B.body_text(doc, para_text)
            else:
                B.body_mixed(doc, text)
        elif kind == "image":
            path, caption = text
            if os.path.exists(path):
                B.add_image(doc, path, caption)
            else:
                print("WARNING: missing image, skipped:", path)
    # New page: separate the LEÇON content from EXERCICES + CORRIGÉ.
    B.page_break(doc)
    B.exercices_heading(doc)
    for header, rest in s["exercices"]:
        for i, line in enumerate(rest.split("\n")):
            if i == 0:
                B.exercise_item(doc, header, line)
            else:
                B.plain_item(doc, line)
    B.total_line(doc, s["total"])
    B.corrige_heading(doc)
    for parts in s["corrige_lines"]:
        B.corrige_mixed(doc, parts)


def build_revexam(doc, s):
    if s.get("cover_image"):
        path, caption = s["cover_image"]
        if os.path.exists(path):
            B.add_image(doc, path, caption)
        else:
            print("WARNING: missing image, skipped:", path)
    B.mixed_para(doc, [(s["sujet_title"], True, B.BLACK)], space_after=8)
    for header, rest in s["exercices"]:
        if header:
            B.para(doc, header, bold=True, space_after=2)
            for line in rest.split("\n"):
                B.plain_item(doc, line)
        else:
            for line in rest.split("\n"):
                B.plain_item(doc, line)
    B.total_line(doc, s["total"])
    B.corrige_heading(doc)
    for parts in s["corrige_lines"]:
        B.corrige_mixed(doc, parts)


def build_seance(doc, s, content):
    B.add_seance_header(
        doc, num=s["num"], total=TOTAL_SEANCES, title=s["title"],
        theme=content["theme"], objectif=content["objectif"], support=content["support"],
        ras_theme=content["ras_theme"], valeurs=content["valeurs"], duree=content["duree"],
    )
    B.add_deroulement_table(doc, content["deroulement"])
    # New page: separate the FICHE DE PRÉPARATION (header + déroulement) from
    # the LEÇON / EXERCICES that follow.
    B.page_break(doc)
    if content.get("kind") in ("revision", "exam"):
        build_revexam(doc, content)
    else:
        build_lecon(doc, content)
    B.page_break(doc)


def _flatten_text(val):
    if isinstance(val, str):
        return val
    if isinstance(val, (list, tuple)):
        parts = []
        for x in val:
            if isinstance(x, (list, tuple)) and len(x) == 2 and isinstance(x[1], bool):
                parts.append(str(x[0]))
            else:
                parts.append(_flatten_text(x))
        return "".join(parts)
    return str(val)


def _all_unit_modules():
    return UNIT_MODULES


def build_annexes(doc):
    import content_annexes as A

    seance_text = {}
    captions = []  # (seance_num, caption) in document order, for Table des illustrations
    gi = 0
    for mod in _all_unit_modules():
        seances = [v for k, v in vars(mod).items() if isinstance(v, dict) and "num" in v]
        seances.sort(key=lambda d: d.get("num", 0))
        for s in seances:
            g = ALL_SEANCES[gi]["num"]
            txt = [s.get("title", "")]
            if s.get("cover_image"):
                captions.append((g, s["cover_image"][1]))
            for item in s.get("lecon", []):
                if isinstance(item, tuple) and len(item) == 2:
                    txt.append(_flatten_text(item[1]))
                    if item[0] == "image":
                        captions.append((g, item[1][1]))
            for row in s.get("deroulement", []):
                for cell in row:
                    txt.append(_flatten_text(cell))
            seance_text[g] = " ".join(txt).lower()
            gi += 1

    index_rows = []
    for term, key, _defn in A.GLOSSARY:
        nums = [n for n in sorted(seance_text) if key.lower() in seance_text[n]]
        index_rows.append((term, nums))
    index_rows.sort(key=lambda r: r[0].lower())

    glossary_sorted = sorted([(t, d) for t, _k, d in A.GLOSSARY], key=lambda r: r[0].lower())

    exam_seances = [(s["unit_roman"], s["num"], s["title"]) for s in ALL_SEANCES if s["kind"] == "exam"]

    B.page_break(doc)
    B.add_annexes_heading(doc)
    bilan_path = os.path.join(HERE, "generated_images", "annexe_bilan.jpg")
    if os.path.exists(bilan_path):
        B.add_image(doc, bilan_path,
                    "Figure A1 — Schéma-bilan : les onze thématiques de Sciences et Technologie T5.",
                    width_cm=13.5)
        B.page_break(doc)

    B.add_annexe_subheading(doc, "GLOSSAIRE", "glossaire")
    B.add_glossary_table(doc, glossary_sorted)
    B.page_break(doc)

    B.add_annexe_subheading(doc, "AUTO-ÉVALUATION", "autoeval")
    B.para(doc, "Coche la case qui correspond à ton niveau après chaque unité.", space_after=8)
    B.add_autoeval_table(doc, A.AUTOEVAL)
    B.page_break(doc)

    B.add_annexe_subheading(doc, "INDEX", "index_annexe")
    B.add_index(doc, index_rows)
    B.page_break(doc)

    B.add_annexe_subheading(doc, "ÉVALUATIONS FORMAT EXAMEN", "evalexamen")
    B.add_exam_links(
        doc, exam_seances,
        "Les sujets d'examen ST T5 de chaque unité sont regroupés ci-dessous pour préparer les "
        "évaluations.",
    )
    B.page_break(doc)

    B.add_annexe_subheading(doc, "TABLE DES ILLUSTRATIONS", "illus")
    B.add_illustrations_table(doc, captions)


def main():
    doc = docx.Document()
    from docx.shared import Inches
    from docx.enum.section import WD_SECTION

    # --- Page 1: full-bleed illustrated cover (zero margins) ---
    cover_section = doc.sections[0]
    cover_section.left_margin = Inches(0)
    cover_section.right_margin = Inches(0)
    cover_section.top_margin = Inches(0)
    cover_section.bottom_margin = Inches(0)
    cover_path = os.path.join(HERE, "generated_images", "cover.jpg")
    if os.path.exists(cover_path):
        B.add_bleed_cover(doc, cover_path)
    else:
        print("WARNING: cover image missing, falling back to text cover:", cover_path)
        B.add_cover(
            doc,
            subject_line="ST",
            level_line="T5",
            programme_line="Programme d'Études — 11 thématiques",
            sub_line="Fiches de préparation · Leçons · Exercices corrigés · Sujets d'examen corrigés",
        )

    # --- Rest of the document: normal narrow margins (0.5 in all around) ---
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    body_section.left_margin = Inches(0.5)
    body_section.right_margin = Inches(0.5)
    body_section.top_margin = Inches(0.5)
    body_section.bottom_margin = Inches(0.5)

    B.add_avant_propos(doc, [
        "Ce manuel de Sciences et Technologie pour la classe de T5 a été élaboré à partir du "
        "Programme d'Études officiel (DCRP, Ministère de l'Éducation Nationale).",
        "Chaque séance est composée d'une fiche de préparation, d'une leçon, d'exercices notés et de leur "
        "corrigé.",
        "Conformément au Programme d'Études, les séances privilégient une approche par démarche "
        "d'investigation, d'observation ou de conception technologique, adaptée au contexte malgache.",
        "Les séances sont regroupées en onze unités, une par thématique du programme. Chaque unité se "
        "termine par une séance de révision puis un sujet d'examen corrigés.",
        "Pour bien utiliser ce manuel, se reporter au « Mode d'emploi » qui suit.",
    ])

    B.add_mode_emploi(doc, [
        "Comment est construite chaque séance ?",
        "I. Révision — questions simples sur la séance précédente, avec la réponse attendue (R.A.).",
        "II. Nouvelle leçon — six sous-étapes : 1. Mise en situation, 2. Présentation, 3. Observation, "
        "4. Analyse, 5. Synthèse, 6. Application.",
        "III. Évaluation — exercices écrits, avec corrigé.",
        "La durée de chaque grande étape n'est pas indiquée : elle est laissée à l'enseignant selon le "
        "contexte de la classe.",
        "Après la fiche de préparation viennent la page LEÇON (le contenu à retenir) puis la section "
        "EXERCICES et son CORRIGÉ.",
    ])

    B.add_global_toc(doc, ALL_SEANCES, UNITS)

    # ---- Build each unit from its content_unitN module ----
    for idx, (roman, title, hours, valeurs, ras, seances_list) in enumerate(UNITS):
        unit_seances = [s for s in ALL_SEANCES if s["unit_roman"] == roman]
        B.add_unit_divider(doc, roman, title, ras, valeurs, unit_seances)

        mod = UNIT_MODULES[idx]
        contents = [v for k, v in vars(mod).items() if isinstance(v, dict) and "num" in v]
        contents.sort(key=lambda d: d["num"])
        if len(contents) != len(unit_seances):
            print(f"WARNING: unit {roman} has {len(contents)} seances in module "
                  f"but {len(unit_seances)} expected in plan.py")
        for s, content in zip(unit_seances, contents):
            build_seance(doc, s, content)

    # ---- ANNEXES ----
    build_annexes(doc)

    doc.save(OUT)
    print("Saved:", OUT)
    print("Paragraphs:", len(doc.paragraphs), "Tables:", len(doc.tables))


if __name__ == "__main__":
    main()
