"""Insert generated images into the T8 docx at their placeholder paragraphs.
Run after build.py. Only inserts images whose spec status == 'done' and whose
corresponding placeholder paragraph is still empty (idempotent / resumable)."""
import json, os, re
from io import BytesIO
import docx
from docx.shared import Inches
from PIL import Image

DOCX_PATH = "../../SVT T8 [PE] Fiche de preparation sujet corriges J-Learn.docx"
SPECS_PATH = "image_specs.json"

def full_caption(spec):
    sid = spec['id']
    m = re.match(r"fig(\d+)(b?)$", sid)
    if m:
        num, b = m.groups()
        num = int(num)
        if b:
            return f"Figure {num}b — {spec['title']}."
        return f"Figure {num} — {spec['title']}."
    m = re.match(r"fig([RE])(\d+)$", sid)
    if m:
        letter, num = m.groups()
        return f"Figure {letter}{num} — {spec['title']}."
    raise ValueError(sid)

def compress_image(path, max_width=1400, quality=82):
    im = Image.open(path).convert("RGB")
    if im.width > max_width:
        ratio = max_width / im.width
        im = im.resize((max_width, int(im.height * ratio)), Image.LANCZOS)
    buf = BytesIO()
    im.save(buf, format="JPEG", quality=quality, optimize=True)
    buf.seek(0)
    return buf

def main():
    specs = json.load(open(SPECS_PATH, encoding="utf-8"))
    caption_map = {full_caption(s): s for s in specs if s['status'] == 'done'}
    print(f"{len(caption_map)} images marked done, looking for placeholders...")

    d = docx.Document(DOCX_PATH)
    paras = d.paragraphs  # body-level only; all placeholders are body-level, not in tables
    inserted = 0
    skipped_no_placeholder = 0
    already_has_image = 0

    for i, p in enumerate(paras):
        text = p.text.strip()
        if text in caption_map:
            spec = caption_map[text]
            if i == 0:
                continue
            placeholder = paras[i - 1]
            # check if placeholder already has a drawing (avoid double insert)
            if placeholder._p.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline'):
                already_has_image += 1
                continue
            if placeholder.text.strip() != "":
                skipped_no_placeholder += 1
                print("  WARNING placeholder not empty for", spec['id'], ":", repr(placeholder.text[:50]))
                continue
            img_path = os.path.join("/home/user/jlearn-pack", spec['file'])
            if not os.path.exists(img_path):
                print("  MISSING FILE for", spec['id'], img_path)
                continue
            buf = compress_image(img_path)
            run = placeholder.add_run()
            run.add_picture(buf, width=Inches(5.5))
            inserted += 1

    d.save(DOCX_PATH)
    print(f"Inserted: {inserted}, already had image: {already_has_image}, "
          f"skipped (no empty placeholder): {skipped_no_placeholder}")

if __name__ == "__main__":
    main()
