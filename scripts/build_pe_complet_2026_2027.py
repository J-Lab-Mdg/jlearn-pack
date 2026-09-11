# -*- coding: utf-8 -*-
"""PE COMPLET 2026-2027 (une seule page-tableau pour les 3 caisses) :
titre « LAMIM-PANDANIANA Tetikasan-tSekoly Ifanekena » ;
colonne LOHARANO NIAVIAN'NY VOLA = CAISSE ECOLE R.P.I (ETAT) / CAISSE ECOLE F.C.E / VOI ;
N° qui se suivent jusqu'à la fin ; ST gras italique sous chaque série.
Données = celles du TeSI v2 (build_pec_tesi_v2.rows).
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_pe_2026_2027 import FONT, fmt, set_cell, header_block
import build_pec_tesi_v2 as TESI
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "pe-2026-2027")
os.makedirs(OUT, exist_ok=True)


def norm_loha(loha, item):
    """Source de fonds affichée dans LOHARANO NIAVIAN'NY VOLA."""
    l = loha.lower()
    if l.startswith("voi") and "rpi" in l:          # série mixte pascoma/sports
        return "VOI" if "pascoma" in item.lower() else "CAISSE ECOLE R.P.I (ETAT)"
    if l.startswith("voi"):
        return "VOI"
    if "fce" in l:
        return "CAISSE ECOLE F.C.E"
    if "rpi" in l:
        return "CAISSE ECOLE R.P.I (ETAT)"
    return loha


def items_of(asa, isa, vidin, tot):
    """retire la ligne d'en-tête (Fividianana :, Fandoavana …) et aligne les nombres."""
    off = len(asa) - len(isa)
    names = [a.lstrip("• ").strip() if a.startswith("•") else a.strip() for a in asa[off:]]
    return list(zip(names, isa, vidin, tot))


doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(9)
sec = doc.sections[0]
sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
sec.top_margin = Cm(1.5); sec.bottom_margin = Cm(1.5)
sec.left_margin = Cm(1.5); sec.right_margin = Cm(1.5)

header_block(doc, "LAMIM-PANDANIANA Tetikasan-tSekoly Ifanekena")

# aplatissement : (lahasa, item, st_si_dernier)
flat = []
for (sehatra, lahasa, fepotoana, series) in TESI.rows:
    for (asa, isa, vidin, tot, stv, loha, mpanat, mpiand) in series:
        its = items_of(asa, isa, vidin, tot)
        for k, it in enumerate(its):
            flat.append((lahasa, it, norm_loha(loha, it[0]), stv if k == len(its) - 1 else None))

n = len(flat)
t = doc.add_table(rows=2 + n, cols=8)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER

set_cell(t.cell(0, 0), "N°", bold=True, align="center")
set_cell(t.cell(0, 1), "LAHASA", bold=True, align="center")
set_cell(t.cell(0, 2), "KARAZANY", bold=True, align="center")
set_cell(t.cell(0, 3), "TETIBIDY", bold=True, align="center")
t.cell(0, 3).merge(t.cell(0, 5))
set_cell(t.cell(0, 6), "LOHARANO\nNIAVIAN’NY\nVOLA", bold=True, align="center")
set_cell(t.cell(0, 7), "FANAMARIHANA", bold=True, align="center")
t.cell(0, 0).merge(t.cell(1, 0))
t.cell(0, 1).merge(t.cell(1, 1))
t.cell(0, 2).merge(t.cell(1, 2))
t.cell(0, 6).merge(t.cell(1, 6))
t.cell(0, 7).merge(t.cell(1, 7))
set_cell(t.cell(1, 3), "Vidin’ny iray", bold=True, align="center")
set_cell(t.cell(1, 4), "Isany", bold=True, align="center")
set_cell(t.cell(1, 5), "Totaly", bold=True, align="center")

num = 0
for r, (lahasa, (kar, isa, vid, tot), loha, stv) in enumerate(flat):
    num += 1
    rr = 2 + r
    set_cell(t.cell(rr, 0), str(num), align="center")
    set_cell(t.cell(rr, 1), lahasa or "", size=8)
    set_cell(t.cell(rr, 2), kar)
    set_cell(t.cell(rr, 3), vid, align="center")
    set_cell(t.cell(rr, 4), isa, align="center")
    if stv is not None:
        set_cell(t.cell(rr, 5), [(tot, False, False), (f"ST : {fmt(stv)}", True, True)], align="center")
    else:
        set_cell(t.cell(rr, 5), tot, align="center")
    set_cell(t.cell(rr, 6), loha, size=8, align="center")
    set_cell(t.cell(rr, 7), "")

# fusion LAHASA par ligne logique (groupes consécutifs identiques)
start = 2
cur = flat[0][0]
for r in range(2, 2 + n + 1):
    la = flat[r - 2][0] if r < 2 + n else None
    if la != cur:
        if r - 1 > start:
            t.cell(start, 1).merge(t.cell(r - 1, 1))
        set_cell(t.cell(start, 1), cur or "", size=8)
        start = r
        cur = la

# fusion LOHARANO par blocs consécutifs identiques
start = 2
cur = flat[0][2]
for r in range(2, 2 + n + 1):
    la = flat[r - 2][2] if r < 2 + n else None
    if la != cur:
        if r - 1 > start:
            t.cell(start, 6).merge(t.cell(r - 1, 6))
        start = r
        cur = la

for txt in [
    "Nofaranana tamin’ny tetibola : enina arivo sy valo alina sy roa hetsy sy sivy tapitrisa Ariary (Ar 9 286 000)",
    "Ka ny 7 795 000 Ar dia avy amin’ny Vondrompiarahamonina (VOI) ary ny 1 491 000 Ar avy amin’ny « Caisse Ecole » (RPI + FCE).",
    "Natao teto Andranomafana, androany faha-04 Septambra 2026",
]:
    p = doc.add_paragraph()
    r = p.add_run(txt)
    r.font.name = FONT
    r.font.size = Pt(10)
    p.paragraph_format.space_before = Pt(8)

ts = doc.add_table(rows=1, cols=3)
from build_pe_2026_2027 import borders_off
borders_off(ts)
for j, txt in enumerate(["NY FILOHA BM FEFFI", "NY SEKRETERAN’NY FEFFI", "MPANDRINDRA NY FB"]):
    set_cell(ts.cell(0, j), txt, bold=True, size=10, align="center")

path = os.path.join(OUT, "PE_COMPLET_2026_2027.docx")
doc.save(path)
print("OK ->", path, "| lignes:", n)
