# -*- coding: utf-8 -*-
"""Génère PEC_TeSI_2026_2027_v2.docx : version corrigée (remarques du chef) + mise en page v2 :
- Isa / Vidin'ny iray / Totalim-bola : UN NOMBRE PAR LIGNE, aligné sur la liste des articles
- ST en GRAS ITALIQUE sous chaque série de totalim-bola
- Bloc résumé (RAM, effectifs, totaux) À DROITE ; titre TETIKASAN/FEFFI/Taom CENTRÉ
- Ankizy voailika : 0 ; date : faha-04 Septambra 2026
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_pe_2026_2027 import FONT, fmt, set_cell, borders_off
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "pe-2026-2027")
os.makedirs(OUT, exist_ok=True)

doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(9)
sec = doc.sections[0]
sec.page_width = Cm(29.7); sec.page_height = Cm(21.0)  # paysage
sec.left_margin = Cm(1.2); sec.right_margin = Cm(1.2)
sec.top_margin = Cm(1.2); sec.bottom_margin = Cm(1.2)

# ---------- en-tête ----------
t0 = doc.add_table(rows=1, cols=2)
borders_off(t0)
left, right = t0.cell(0, 0), t0.cell(0, 1)
left.text = ""
for i, txt in enumerate([
    "MINISTERAN’NY FANABEAZAM-PIRENENA", "DREN : ALAOTRA MANGORO", "CISCO : ANOSIBE AN’ALA",
    "ZAP : TRATRAMARINA", "EPP : ANDRANOMAFANA", "MARIPANGO : 506 090 012"]):
    p = left.paragraphs[0] if i == 0 else left.add_paragraph()
    r = p.add_run(txt); r.font.name = FONT; r.font.size = Pt(10); r.font.bold = True
    p.paragraph_format.space_after = Pt(1)
right.text = ""
centre = ["TETIKASAN-TSEKOLY IFANEKENA", "FEFFI : ANDRANOMAFANA", "Taom-pianarana 2026 – 2027"]
adroite = [
    "Totalin’ny vola haloan’ny RAM : 5 029 000 Ar",
    "Vola haloan’ny RAM iray : 100 580 Ar",
    "Isan’ny Mpianatra tamin’ny taon-dasa : V 68   L 61   T 129",
    "Isan’ny Mpianatra amin’ny taona diavina : V 82   L 63   T 145",
    "Totalin’ny lahasa voalahatra : 12",
    "Totalin’ny lahasa voalahatra ho an’ny VOI : 06",
    "Tetibidy novinavinaina : 9 286 000 Ar",
    "Tetibidy novinavinain’ny VOI : 7 795 000 Ar"]
for i, txt in enumerate(centre + adroite):
    p = right.paragraphs[0] if i == 0 else right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i < 3 else WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(txt); r.font.name = FONT; r.font.size = Pt(9); r.font.bold = (i < 3)
    p.paragraph_format.space_after = Pt(1)

# ---------- données : isa/vidin/tot = listes (un nombre par ligne) ----------
SIG_RPI = "Ny Filoha, Ny Tale, Ny Trésorier"
SIG_FCE = "–"
SUP_RPI = "Mpanara-maso ny kaonty, Tale, Membre FEFFI"
SUP_FCE = "Membre FEFFI, Tale"
SUP_VOI = "Filohan’ny FEFFI, Tale"
SUP_BIRAO = "Ny Filohan’ny FEFFI, Ny Mpitam-bola, Ny Mpitam-tsoratra"

rows = []
rows.append(("FOTOANA IANARANA", "Fanatanterahana ny fianarana (Fianarana sy ny Ora Fanarenana « OF »)",
             "Septambra 2026 – Jona 2027", [
    (["Fividianana :", "• Craie couleur RAFAFA", "• Craie blanche RAFAFA", "• Eponge", "• Isorel",
      "• Ody parasy sachets", "• Siflet vy"],
     ["10", "25", "21", "2", "3", "6"], ["6 000", "5 000", "700", "58 000", "3 000", "1 800"],
     ["60 000", "125 000", "14 700", "116 000", "9 000", "10 800"], 335500,
     "Caisse école : RPI (Etat)", SIG_RPI, SUP_RPI),
    (["• Ardoisine GM", "• Agrafe 4/6 boîte", "• Craie Rafafa couleur", "• Craie blanche RAFAFA"],
     ["2", "6", "10", "10"], ["12 000", "1 800", "6 000", "5 000"],
     ["24 000", "10 800", "60 000", "50 000"], 144800, "Caisse école : FCE", SIG_FCE, SUP_FCE),
]))

rows.append((None, "Fitiliana ny mpianatra ahiana hitsoaka andaharana (2026-2027)", "Isam-bolana", [
    (["Fanaovana lisitra"], ["N/A"], ["N/A"], ["N/A"], None, "N/A", "Tale Mpampianatra",
     "Filohan’ny FEFFI, Tale"),
]))

rows.append((None, "VAD any an-tokantranon’ny mpianatra ahiana hitsoaka andaharana (2026-2027)", "Isam-bolana", [
    (["Fitiliana / Fandaminana sy Fandravonana VAD"], ["N/A"], ["N/A"], ["N/A"], None, "N/A",
     "Membre FEFFI", "Président, Tale"),
]))

rows.append(("TONTOLO IANARANA / FITAOVANA", "Fanamboarana trano mpampianatra (logement)",
             "Septambra 2026 – Jona 2027", [
    (["Fividianana :", "• Ciment", "• Pointe 5,6,7,8", "• Angady PM"],
     ["2", "2", "1"], ["52 000", "9 000", "5 000"], ["104 000", "18 000", "5 000"], 127000,
     "Caisse école : RPI (Etat)", "Ny Tale, Ny Filoha, Ny Trésorier",
     "Komity ray aman-dreny, Membre FEFFI"),
    (["• Ciment", "• Angady GM", "• Angady PM"], ["1", "1", "2"], ["52 000", "16 000", "5 000"],
     ["52 000", "16 000", "10 000"], 78000, "Caisse école : FCE", SIG_FCE, "–"),
    (["• Andry", "• Chevron", "• Tôle 2,50", "• Pointe tôle"], ["60", "45", "30", "3"],
     ["6 000", "6 000", "40 000", "9 000"], ["360 000", "270 000", "1 200 000", "27 000"], 1857000,
     "VOI", "VOI", "VOI"),
]))

rows.append((None, "Fividianana fitaovana ho an’ny mpampianatra (côte bureau)",
             "Septambra – Desambra 2026", [
    (["Fividianana :", "• Colle bureau", "• Wood Fix 500 g", "• Cahier GF 200p lux", "• Cahier PF 200p lux",
      "• Chemise à dossier", "• Ciseaux PM", "• Enveloppe A4 Kraft", "• Marker Dollar 6 BL/ER",
      "• Papier Velin ram", "• Pinceau N°3", "• Rateau", "• Règle pour TN", "• Stylo Digno R/N/B (6×3)",
      "• Stylo Schneider Bleu", "• Stylo Schneider Rouge", "• Stylo Schneider Noir", "• Stylo Schneider Vert",
      "• Scotch GM", "• Tadin-tsaina", "• Emballage blanc", "• Badge", "• Porte feuille", "• Registre d’appel GM"],
     ["6", "1", "14", "12", "60", "6", "18", "12", "2", "2", "1", "1", "18", "18", "18", "18", "6", "6", "24", "36", "1", "12", "7"],
     ["1 600", "9 000", "5 800", "3 600", "600", "1 800", "500", "3 000", "25 000", "2 500", "8 500", "10 500", "700", "1 000", "1 000", "1 000", "1 000", "6 000", "400", "400", "2 000", "300", "3 500"],
     ["9 600", "9 000", "81 200", "43 200", "36 000", "10 800", "9 000", "36 000", "50 000", "5 000", "8 500", "10 500", "12 600", "18 000", "18 000", "18 000", "6 000", "36 000", "9 600", "14 400", "2 000", "3 600", "24 500"],
     471500, "Caisse école : RPI (Etat)", SIG_RPI, SUP_BIRAO),
    (["• Blanco pen", "• Cahier de transmission", "• Cahier de départ", "• Cahier d’arrivée",
      "• Décamètre à 20 m", "• Masking Tape", "• Nappe de table lamba/m", "• Papier Velin ram",
      "• Enveloppe blanc PM"],
     ["6", "1", "1", "1", "1", "6", "1", "1", "12"],
     ["2 000", "11 000", "8 000", "8 000", "18 000", "3 000", "9 000", "25 000", "100"],
     ["12 000", "11 000", "8 000", "8 000", "18 000", "18 000", "9 000", "25 000", "1 200"], 110200,
     "Caisse école : FCE", SIG_FCE, SUP_BIRAO),
]))

rows.append(("KALITAOM-PAMPIANARANA SY FIANARANA", "Fandraisana mpampianatra",
             "Septambra 2026 – Jona 2027", [
    (["Fandoavana :", "• Karamana mpampianatra"], ["4"], ["1 250 000"], ["5 000 000"], 5000000,
     "VOI", "VOI", SUP_VOI),
]))

rows.append((None, "Famonjena Journées Pédagogiques (JP) sy Sports scolaires / Hetsika",
             "Septambra 2026 – Jona 2027", [
    (["Fandoavana tsipaipaika :", "• Journées Pédagogiques", "• Sports scolaires", "• Hetsika « vokatra »"],
     ["5", "5", "2"], ["72 000", "50 000", "35 000"], ["360 000", "250 000", "70 000"], 680000,
     "VOI", "VOI", SUP_VOI),
]))

rows.append((None, "Fandoavana ny fiantohana « Pascoma » sy ny Sports scolaires",
             "Septambra 2026 – Janoary 2027", [
    (["Fandoavana :", "• Pascoma", "• Sports scolaire"], ["140", "140"], ["200", "500"],
     ["28 000", "70 000"], None, "VOI ST : 28 000   RPI ST : 70 000", "VOI / FEFFI", SUP_VOI),
]))

rows.append((None, "Fampisaafoanana an-tsekoly (Fambolena)", "Oktobra 2026 – Desambra 2026", [
    (["Fambolena :", "• Mangahazo (tokony)", "• Siramamy (kg)"], ["60", "20"], ["1 000", "6 000"],
     ["60 000", "120 000"], 180000, "VOI", "VOI", SUP_VOI),
]))

rows.append(("FAMPANDEHANANA NY FEFFI SY VFF", "Fampitaovana ny birao", "Septambra 2026", [
    (["Fividianana :", "• Gony 250 kg"], ["1"], ["4 000"], ["4 000"], 4000, "Caisse école : FCE",
     "FEFFI", SUP_BIRAO),
]))

rows.append((None, "Fampandraharana ny VFF", "Septambra 2026 – Janoary 2027", [
    (["Hetsika VFF"], ["1"], ["50 000"], ["50 000"], 50000, "VOI", "VOI", SUP_VOI),
]))

rows.append((None, "Fivezivezen’ny mpiambana ao amin’ny BM FEFFI", "Septambra 2026", [
    (["• Saran-dalana (FCE)"], ["3"], ["20 000"], ["60 000"], 60000, "Caisse école : FCE", "FEFFI",
     SUP_BIRAO),
    (["• Saran-dalana (RPI)"], ["3"], ["30 000"], ["90 000"], 90000, "Caisse école : RPI", "FEFFI",
     SUP_BIRAO),
]))

# ---------- tableau ----------
COLS = 10
n_series = sum(len(r[3]) for r in rows)
t = doc.add_table(rows=2 + n_series, cols=COLS)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER

for j, htxt in [(0, "SEHATRA"), (1, "LAHASA"), (2, "FE-POTOANA"), (3, "ASA"),
                (7, "LOHARANOM-\nBOLA"), (8, "MPANATA-\nNTERAKA"), (9, "MPIANDRAIKITRA\nNY FANJOHIANA")]:
    set_cell(t.cell(0, j), htxt, bold=True, size=8, align="center")
    t.cell(0, j).merge(t.cell(1, j))
set_cell(t.cell(0, 4), "TETIBIDY", bold=True, size=8, align="center")
t.cell(0, 4).merge(t.cell(0, 6))
set_cell(t.cell(1, 4), "Isa", bold=True, size=8, align="center")
set_cell(t.cell(1, 5), "Vidin’ny iray", bold=True, size=8, align="center")
set_cell(t.cell(1, 6), "Totalim-bola", bold=True, size=8, align="center")

r0 = 2
for (sehatra, lahasa, fepotoana, series) in rows:
    start = r0
    for (asa, isa, vidin, tot, stv, loha, mpanat, mpiand) in series:
        # une ligne par article ; Isa/Vidin/Totalim alignés ligne à ligne
        n = max(len(asa), len(isa))
        asa_l = list(asa) + [""] * (n - len(asa))
        isa_l = [""] * (n - len(isa)) + list(isa)          # nombres en bas, face aux articles
        vid_l = [""] * (n - len(vidin)) + list(vidin)
        tot_l = [""] * (n - len(tot)) + list(tot)
        lines = [(a, False, False) for a in asa_l]
        for k in range(n):
            pass
        set_cell(t.cell(r0, 3), [(a, False, False) for a in asa_l])
        set_cell(t.cell(r0, 4), [(x, False, False) for x in isa_l], align="center")
        set_cell(t.cell(r0, 5), [(x, False, False) for x in vid_l], align="center")
        if stv is not None:
            tot_l.append(f"ST : {fmt(stv)}")
            st_idx = len(tot_l) - 1
            set_cell(t.cell(r0, 6), [(x, (i == st_idx), (i == st_idx)) for i, x in enumerate(tot_l)],
                     align="center")
        else:
            set_cell(t.cell(r0, 6), [(x, False, False) for x in tot_l], align="center")
        set_cell(t.cell(r0, 7), loha, size=8, align="center")
        set_cell(t.cell(r0, 8), mpanat, size=8, align="center")
        set_cell(t.cell(r0, 9), mpiand, size=8, align="center")
        r0 += 1
    end = r0 - 1
    if end > start:
        for j in (1, 2):
            t.cell(start, j).merge(t.cell(end, j))
    set_cell(t.cell(start, 1), lahasa or "", size=8)
    set_cell(t.cell(start, 2), fepotoana or "", size=8)

# fusion SEHATRA par blocs
blocks = []
r = 2
cur = None
bs = None
for (sehatra, lahasa, fepotoana, series) in rows:
    if sehatra is not None:
        if bs is not None:
            blocks.append((bs, r - 1, cur))
        bs = r
        cur = sehatra
    r += len(series)
blocks.append((bs, r - 1, cur))
for (a, b, lab) in blocks:
    if b > a:
        t.cell(a, 0).merge(t.cell(b, 0))
    set_cell(t.cell(a, 0), lab or "", bold=True, size=8)

# ---------- pied ----------
for txt in [
    "Totalin’ny lahasa voalahatra : 12 ka ny lahasa 06 dia avy amin’ny entimanan’ny Vondrompiarahamonina (VOI).  Tetibidy novinavinaina : 9 286 000 Ar, ka ny 7 795 000 Ar dia avy amin’ny Vondrompiarahamonina ary ny 1 491 000 Ar avy amin’ny « Caisse Ecole ».",
    "Fitambaran’ny isan’ny Mpampianatra novinavinain’ny Vondrompiarahamonina horaisina an-tanana : 04",
    "Totalin’ny Ora voatokana hanaovana fampianarana fanarenana isan-kilasy mandritra ny taona : 108 Ora (kilasy voakasik’izany T2, T3, T4, T5)",
    "Ankizy voailika niditra an-tsekoly : 0",
    "Natao teto Andranomafana, androany faha-04 Septambra 2026",
]:
    p = doc.add_paragraph()
    r = p.add_run(txt); r.font.name = FONT; r.font.size = Pt(9); r.font.bold = txt.startswith("Totalin’ny lahasa")
    p.paragraph_format.space_before = Pt(6)

ts = doc.add_table(rows=1, cols=3)
borders_off(ts)
for j, txt in enumerate(["NY FILOHA BM FEFFI", "NY SEKRETERAN’NY FEFFI", "MPANDRINDRA NY FB"]):
    set_cell(ts.cell(0, j), txt, bold=True, size=10, align="center")

path = os.path.join(OUT, "PEC_TeSI_2026_2027_v2.docx")
doc.save(path)
print("OK ->", path)
