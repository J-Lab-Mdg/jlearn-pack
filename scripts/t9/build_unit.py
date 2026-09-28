# -*- coding: utf-8 -*-
"""Transplant Manuel-SVT-T9-complet content into the official T9 Fiche de
preparation docx, seance by seance, unit by unit.

Usage: python3 build_unit.py <first_seance> <last_seance> <src_docx> <out_docx>
"""
import copy
import os
import pickle
import re
import sys

import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
from docx.table import Table

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))

Wp = qn('w:p')
Wtbl = qn('w:tbl')
Wt = qn('w:t')
Wr = qn('w:r')

# ---- RAS / Valeurs authoritative mapping (from PE_T9) ----
UNIT_VALEURS = {
    1: "autonomie, esprit de créativité",
    2: "respect de toute vie, culture de l’excellence",
    3: "connaissance de soi, altruisme",
    4: "estime de soi, responsabilité",
    5: "responsabilité, sens du bien commun",
}
UNIT_THEME = {
    1: "Alimentation de l’homme",
    2: "Organisation des êtres vivants",
    3: "Santé et bien-être",
    4: "Reproduction humaine",
    5: "Géologie",
}
UNIT_ROMAN = {1: 'I', 2: 'II', 3: 'III', 4: 'IV', 5: 'V'}

UNIT_RAS_SUMMARY = {
    1: "Proposer des menus équilibrés ; proposer des menus ou des régimes d’appoint "
       "en cas de malnutrition (carence ou excès)",
    2: "Proposer des soins appropriés à un type de plante ; mettre en relation les "
       "techniques d’amélioration de l’élevage et l’espèce concernée",
    3: "Participer à la lutte contre la prise des substances psychoactives",
    4: "Participer à la lutte contre la fistule obstétricale et le cancer du col de l’utérus",
    5: "Mettre en relation les caractéristiques et l’utilisation des pierres gemmes ; "
       "analyser les impacts environnementaux de leur exploitation",
}

RAS_BY_SEANCE = {}
for n in range(1, 9):
    RAS_BY_SEANCE[n] = (1, "Proposer des menus équilibrés")
for n in (9, 10):
    RAS_BY_SEANCE[n] = (1, "Proposer des menus ou des régimes d’appoint pour un individu dans "
                            "une situation de malnutrition (carence ou excès)")
for n in range(13, 19):
    RAS_BY_SEANCE[n] = (2, "Proposer des soins appropriés à un type de plante")
for n in range(19, 23):
    RAS_BY_SEANCE[n] = (2, "Mettre en relation différentes techniques d’amélioration de "
                            "l’élevage à chaque espèce connue")
for n in range(25, 34):
    RAS_BY_SEANCE[n] = (3, "Participer à la lutte contre la prise des substances psychoactives")
for n in range(36, 42):
    RAS_BY_SEANCE[n] = (4, "Participer à la lutte contre la fistule obstétricale et le cancer "
                            "du col de l’utérus")
for n in range(44, 48):
    RAS_BY_SEANCE[n] = (5, "Mettre en relation les caractéristiques et l’utilisation des pierres gemmes")
for n in (48, 49):
    RAS_BY_SEANCE[n] = (5, "Analyser les impacts environnementaux associés à l’exploitation "
                            "des pierres gemmes")

REVISION_EXAM = {11: 1, 12: 1, 23: 2, 24: 2, 34: 3, 35: 3, 42: 4, 43: 4, 50: 5, 51: 5}


def elem_text(el):
    return ''.join(t.text or '' for t in el.findall('.//' + Wt))


def load_manual():
    with open(os.path.join(HERE, 'manual_seances.pkl'), 'rb') as f:
        return pickle.load(f)


def find_body_children(doc):
    return list(doc.element.body)


def find_seance_marker(children, num):
    pat = re.compile(r'^SÉANCE\s+' + str(num) + r'\s*/\s*51\s*$')
    for el in children:
        if el.tag == Wp and pat.match(elem_text(el).strip()):
            return el
    raise RuntimeError(f"seance marker {num} not found")


def next_sibling_tables(anchor, n=1):
    out = []
    cur = anchor
    while cur is not None and len(out) < n:
        cur = cur.getnext()
        if cur is not None and cur.tag == Wtbl:
            out.append(cur)
    return out


def set_labeled_paragraph(cell, label_prefix, new_value):
    for p in cell.paragraphs:
        full = p.text
        if full.startswith(label_prefix):
            runs = p.runs
            if not runs:
                continue
            if len(runs) >= 2:
                runs[1].text = new_value
                for extra in runs[2:]:
                    extra.text = ''
            else:
                runs[0].text = label_prefix + new_value
            return True
    return False


def rebuild_deroulement_table(table, data_rows, template_data_tr, template_merged_tr):
    tbl_el = table._tbl
    trs = tbl_el.findall(qn('w:tr'))
    for tr in trs[2:]:
        tbl_el.remove(tr)

    insert_after = trs[1]
    for row in data_rows:
        is_separator = len(set(x.strip() for x in row)) == 1 and row[0].strip() != ''
        if is_separator:
            new_tr = copy.deepcopy(template_merged_tr)
            tcs = new_tr.findall(qn('w:tc'))
            t = tcs[0].find('.//' + Wt)
            if t is None:
                p = tcs[0].find(Wp)
                r = OxmlElement('w:r')
                t = OxmlElement('w:t')
                r.append(t)
                p.append(r)
            t.text = row[0]
            for extra_t in tcs[0].findall('.//' + Wt)[1:]:
                extra_t.text = ''
        else:
            new_tr = copy.deepcopy(template_data_tr)
            tcs = new_tr.findall(qn('w:tc'))
            for tc, value in zip(tcs, row):
                lines = value.split('\n') if value else ['']
                paras = tc.findall(Wp)
                n = min(len(lines), len(paras))
                for i in range(n):
                    runs = paras[i].findall(Wr)
                    if runs:
                        t = runs[0].find(Wt)
                        if t is None:
                            t = OxmlElement('w:t')
                            t.set(qn('xml:space'), 'preserve')
                            runs[0].append(t)
                        t.text = lines[i]
                        for r in runs[1:]:
                            tt = r.find(Wt)
                            if tt is not None:
                                tt.text = ''
                    else:
                        r = OxmlElement('w:r')
                        t = OxmlElement('w:t')
                        t.set(qn('xml:space'), 'preserve')
                        t.text = lines[i]
                        r.append(t)
                        paras[i].append(r)
                if len(lines) > len(paras) and paras:
                    last_p = paras[-1]
                    for extra in lines[len(paras):]:
                        newp = copy.deepcopy(last_p)
                        rs = newp.findall(Wr)
                        for i2, r in enumerate(rs):
                            t = r.find(Wt)
                            if t is not None:
                                t.text = extra if i2 == 0 else ''
                        last_p.addnext(newp)
                        last_p = newp
                if len(lines) < len(paras):
                    for p in paras[len(lines):]:
                        for r in p.findall(Wr):
                            t = r.find(Wt)
                            if t is not None:
                                t.text = ''
        insert_after.addnext(new_tr)
        insert_after = new_tr


def remove_zone(start_exclusive, end_exclusive):
    cur = start_exclusive.getnext()
    while cur is not None and cur is not end_exclusive:
        nxt = cur.getnext()
        cur.getparent().remove(cur)
        cur = nxt


def insert_after_anchor(anchor, xml_blobs):
    insert_point = anchor
    for blob in xml_blobs:
        new_el = parse_xml(blob)
        insert_point.addnext(new_el)
        insert_point = new_el
    return insert_point


def find_next_marker_or_none(after_el, patterns):
    cur = after_el.getnext()
    while cur is not None:
        if cur.tag == Wp:
            txt = elem_text(cur).strip()
            for pat in patterns:
                if pat.match(txt):
                    return cur
        cur = cur.getnext()
    return None


def make_note_paragraph(template_p_el, text):
    new_p = copy.deepcopy(template_p_el)
    runs = new_p.findall('.//' + Wr)
    if runs:
        t = runs[0].find(Wt)
        if t is None:
            t = OxmlElement('w:t')
            runs[0].append(t)
        t.text = text
        for r in runs[1:]:
            tt = r.find(Wt)
            if tt is not None:
                tt.text = ''
    return new_p


def process_normal_lesson(doc, num, manual_seance, tpl):
    children = find_body_children(doc)
    marker = find_seance_marker(children, num)
    header_table_el, deroulement_table_el = next_sibling_tables(marker, 2)
    header_table = Table(header_table_el, doc)
    deroulement_table = Table(deroulement_table_el, doc)

    left, right = header_table.rows[0].cells[0], header_table.rows[0].cells[1]
    if manual_seance['objectif']:
        set_labeled_paragraph(left, 'Objectif spécifique : ', manual_seance['objectif'])
    if manual_seance['support']:
        set_labeled_paragraph(left, 'Support et matériel : ', manual_seance['support'])
    unit, ras_text = RAS_BY_SEANCE[num]
    set_labeled_paragraph(right, 'Thématique / RAS : ', ras_text)
    set_labeled_paragraph(right, 'Valeurs à véhiculer : ', UNIT_VALEURS[unit])

    if manual_seance['deroulement_rows']:
        rebuild_deroulement_table(deroulement_table, manual_seance['deroulement_rows'],
                                   tpl['data_tr'], tpl['merged_tr'])

    cur = deroulement_table_el.getnext()
    title_el = None
    while cur is not None:
        if cur.tag == Wp and elem_text(cur).strip():
            title_el = cur
            break
        cur = cur.getnext()
    assert title_el is not None, f"title not found for seance {num}"

    exercices_pat = re.compile(r'^EXERCICES$')
    corrige_pat = re.compile(r'^CORRIGÉ$')
    total_pat = re.compile(r'^TOTAL\s*:')

    exercices_el = None
    corrige_el = None
    total_el = None
    cur = title_el.getnext()
    while cur is not None:
        if cur.tag == Wp:
            txt = elem_text(cur).strip()
            if exercices_pat.match(txt) and exercices_el is None:
                exercices_el = cur
            elif corrige_pat.match(txt) and corrige_el is None:
                corrige_el = cur
            elif total_pat.match(txt):
                total_el = cur
        if corrige_el is not None and cur is not corrige_el:
            txt = elem_text(cur).strip() if cur.tag == Wp else ''
            if re.match(r'^SÉANCE\s+\d+\s*/\s*51\s*$', txt):
                break
        cur = cur.getnext()

    assert exercices_el is not None, f"EXERCICES marker not found for seance {num}"
    assert corrige_el is not None, f"CORRIGÉ marker not found for seance {num}"
    assert total_el is not None, f"TOTAL line not found for seance {num}"

    end_marker = find_next_marker_or_none(corrige_el,
                                           [re.compile(r'^SÉANCE\s+\d+\s*/\s*51\s*$'),
                                            re.compile(r'^UNITÉ\s+[IVX]+')])

    # LECON zone: keep existing images in place, drop everything else, insert
    # manual's lecon content (plus an optional teacher's note for sensitive topics).
    kept_images = []
    cur = title_el.getnext()
    while cur is not None and cur is not exercices_el:
        nxt = cur.getnext()
        if cur.tag == Wp and cur.find('.//' + qn('w:drawing')) is not None:
            kept_images.append(cur)
        else:
            cur.getparent().remove(cur)
        cur = nxt
    insert_point = kept_images[-1] if kept_images else title_el
    if manual_seance.get('note_enseignant'):
        note_p = make_note_paragraph(tpl['note_p_template'], manual_seance['note_enseignant'])
        insert_point.addnext(note_p)
        insert_point = note_p
    if manual_seance['lecon_elements']:
        insert_point = insert_after_anchor(insert_point, manual_seance['lecon_elements'])

    remove_zone(exercices_el, total_el)
    if manual_seance['exercices_elements']:
        insert_after_anchor(exercices_el, manual_seance['exercices_elements'])

    if manual_seance['total_text']:
        runs = total_el.findall('.//' + Wr)
        if runs:
            t = runs[0].find(Wt)
            if t is not None:
                t.text = manual_seance['total_text']
            for r in runs[1:]:
                tt = r.find(Wt)
                if tt is not None:
                    tt.text = ''

    remove_zone(corrige_el, end_marker)
    if manual_seance['corrige_elements']:
        insert_after_anchor(corrige_el, manual_seance['corrige_elements'])

    print(f"  seance {num}: OK (kept {len(kept_images)} images)")


def process_revision_exam(doc, num, manual_seance, tpl):
    children = find_body_children(doc)
    marker = find_seance_marker(children, num)
    title_el = marker.getnext()
    assert title_el is not None and title_el.tag == Wp

    end_marker = find_next_marker_or_none(title_el,
                                           [re.compile(r'^SÉANCE\s+\d+\s*/\s*51\s*$'),
                                            re.compile(r'^UNITÉ\s+[IVX]+')])

    kept = []
    to_remove = []
    cur = title_el.getnext()
    while cur is not None and cur is not end_marker:
        nxt = cur.getnext()
        if cur.tag == Wp and cur.find('.//' + qn('w:drawing')) is not None:
            kept.append(cur)
            if nxt is not None and nxt is not end_marker and nxt.tag == Wp and elem_text(nxt).strip():
                kept.append(nxt)
                nxt = nxt.getnext()
        else:
            to_remove.append(cur)
        cur = nxt
    for el in to_remove:
        el.getparent().remove(el)

    insert_point = kept[-1] if kept else title_el

    if manual_seance.get('note_enseignant'):
        note_p = make_note_paragraph(tpl['note_p_template'], manual_seance['note_enseignant'])
        insert_point.addnext(note_p)
        insert_point = note_p

    new_objectif = copy.deepcopy(tpl['objectif_p'])
    runs = new_objectif.findall('.//' + Wr)
    if len(runs) >= 2:
        runs[1].find(Wt).text = manual_seance['objectif'] or ''
    insert_point.addnext(new_objectif)
    insert_point = new_objectif

    new_support = copy.deepcopy(tpl['support_p'])
    runs = new_support.findall('.//' + Wr)
    if len(runs) >= 2:
        runs[1].find(Wt).text = manual_seance['support'] or ''
    insert_point.addnext(new_support)
    insert_point = new_support

    new_table_el = copy.deepcopy(tpl['deroulement_table'])
    insert_point.addnext(new_table_el)
    insert_point = new_table_el
    new_table = Table(new_table_el, doc)
    if manual_seance['deroulement_rows']:
        rebuild_deroulement_table(new_table, manual_seance['deroulement_rows'],
                                   tpl['data_tr'], tpl['merged_tr'])

    new_exercices_head = copy.deepcopy(tpl['exercices_p'])
    insert_point.addnext(new_exercices_head)
    insert_point = new_exercices_head
    if manual_seance['exercices_elements']:
        insert_point = insert_after_anchor(insert_point, manual_seance['exercices_elements'])

    if manual_seance['total_text']:
        new_total = copy.deepcopy(tpl['total_p'])
        runs = new_total.findall('.//' + Wr)
        if runs:
            t = runs[0].find(Wt)
            if t is not None:
                t.text = manual_seance['total_text']
            for r in runs[1:]:
                tt = r.find(Wt)
                if tt is not None:
                    tt.text = ''
        insert_point.addnext(new_total)
        insert_point = new_total

    new_corrige_head = copy.deepcopy(tpl['corrige_p'])
    insert_point.addnext(new_corrige_head)
    insert_point = new_corrige_head
    if manual_seance['corrige_elements']:
        insert_point = insert_after_anchor(insert_point, manual_seance['corrige_elements'])

    print(f"  seance {num} (revision/exam): OK (kept {len(kept)} elements: images+captions)")


def fix_unit_divider(doc, unit_num):
    children = find_body_children(doc)
    roman = UNIT_ROMAN[unit_num]
    pat = re.compile(r'^UNITÉ\s+' + roman + r'\s*—')
    divider = None
    for el in children:
        if el.tag == Wp and pat.match(elem_text(el).strip()):
            nxt = el.getnext()
            if nxt is not None and 'apprentissage' in elem_text(nxt):
                divider = el
                break
    if divider is None:
        print(f"  [!] unit {unit_num} divider not found")
        return
    ras_p = divider.getnext()
    valeurs_p = ras_p.getnext() if ras_p is not None else None
    if ras_p is not None and 'apprentissage' in elem_text(ras_p):
        runs = ras_p.findall('.//' + Wr)
        if len(runs) >= 2:
            runs[1].find(Wt).text = UNIT_RAS_SUMMARY[unit_num]
            for r in runs[2:]:
                t = r.find(Wt)
                if t is not None:
                    t.text = ''
        elif len(runs) == 1:
            runs[0].find(Wt).text = ("Résultat d'apprentissage spécifique (PE T9) : "
                                      + UNIT_RAS_SUMMARY[unit_num])
    if valeurs_p is not None and 'Valeurs' in elem_text(valeurs_p):
        runs = valeurs_p.findall('.//' + Wr)
        if len(runs) >= 2:
            runs[1].find(Wt).text = UNIT_VALEURS[unit_num]
            for r in runs[2:]:
                t = r.find(Wt)
                if t is not None:
                    t.text = ''
        elif len(runs) == 1:
            runs[0].find(Wt).text = "Valeurs à véhiculer : " + UNIT_VALEURS[unit_num]
    print(f"  unit {unit_num} divider: OK")


def main():
    first = int(sys.argv[1])
    last = int(sys.argv[2])
    src = sys.argv[3]
    out = sys.argv[4]

    manual = load_manual()
    doc = docx.Document(src)

    children0 = find_body_children(doc)
    marker1 = find_seance_marker(children0, 1)
    header1_el, der1_el = next_sibling_tables(marker1, 2)
    header1 = Table(header1_el, doc)
    der1_trs = der1_el.findall(qn('w:tr'))
    template_data_tr = copy.deepcopy(der1_trs[2])
    template_merged_tr = copy.deepcopy(der1_trs[3])
    template_deroulement_table = copy.deepcopy(der1_el)

    left_cell = header1.rows[0].cells[0]
    objectif_p = support_p = None
    for p in left_cell.paragraphs:
        if p.text.startswith('Objectif spécifique'):
            objectif_p = copy.deepcopy(p._p)
        elif p.text.startswith('Support et matériel'):
            support_p = copy.deepcopy(p._p)

    exercices_p = total_p = corrige_p = None
    cur = der1_el.getnext()
    seance2_marker = find_seance_marker(children0, 2)
    while cur is not None and cur is not seance2_marker:
        if cur.tag == Wp:
            txt = elem_text(cur).strip()
            if txt == 'EXERCICES' and exercices_p is None:
                exercices_p = copy.deepcopy(cur)
            elif txt.startswith('TOTAL :') and total_p is None:
                total_p = copy.deepcopy(cur)
            elif txt == 'CORRIGÉ' and corrige_p is None:
                corrige_p = copy.deepcopy(cur)
        cur = cur.getnext()

    # simple bold-label paragraph template reused for the optional "note a
    # l'enseignant" callouts on sensitive-topic seances
    note_p_template = copy.deepcopy(objectif_p)

    tpl = {
        'data_tr': template_data_tr,
        'merged_tr': template_merged_tr,
        'deroulement_table': template_deroulement_table,
        'objectif_p': objectif_p,
        'support_p': support_p,
        'exercices_p': exercices_p,
        'total_p': total_p,
        'corrige_p': corrige_p,
        'note_p_template': note_p_template,
    }
    assert all(v is not None for v in tpl.values()), "missing a template element"

    units_fixed = set()
    for num in range(first, last + 1):
        unit_of_num = RAS_BY_SEANCE.get(num, (REVISION_EXAM.get(num),))[0]
        if unit_of_num and unit_of_num not in units_fixed:
            fix_unit_divider(doc, unit_of_num)
            units_fixed.add(unit_of_num)
        if num in REVISION_EXAM:
            process_revision_exam(doc, num, manual[num], tpl)
        else:
            process_normal_lesson(doc, num, manual[num], tpl)

    doc.save(out)
    print("Saved:", out)


if __name__ == '__main__':
    main()
