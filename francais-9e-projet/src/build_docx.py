# build_docx.py — régénère les manuels PE/PS en .docx via python-docx (toujours lisible par Word).
import json, os, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "..", "assets")
ROSE = RGBColor(0xE9, 0x70, 0x4F); ROUGE = RGBColor(0xC0, 0x00, 0x00)
VERT = RGBColor(0x1E, 0x7B, 0x34); BLEU = RGBColor(0x1F, 0x4E, 0x79)

MODE = sys.argv[1] if len(sys.argv) > 1 else "PE"
data = json.load(open(os.path.join(HERE, "..", "output", "dump_" + MODE + ".json"), encoding="utf-8"))

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)

def T(v):
    return v if isinstance(v, str) else " ".join(v)

def setfont(run, color=None, bold=False, italic=False, size=None):
    run.font.name = "Times New Roman"; run.bold = bold; run.italic = italic
    if color: run.font.color.rgb = color
    if size: run.font.size = Pt(size)

def h(text, level, color=None):
    p = doc.add_heading(level=level)
    r = p.add_run(text); setfont(r, color or (ROUGE if level == 1 else None), bold=True)
    return p

def para(text="", color=None, bold=False, italic=False, size=None, align=None):
    p = doc.add_paragraph()
    if align: p.alignment = align
    if text:
        r = p.add_run(text); setfont(r, color, bold, italic, size)
    return p

# Couverture
para("COLLECTION J-LEARN", ROUGE, True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
para("MANUEL DE FRANÇAIS", ROUGE, True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER)
para("CLASSE DE 9ÈME (T3 / CE)", None, True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Programme d'Études (PE)" if MODE == "PE" else "Programme Scolaire (PS)", None, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

for th in data["themes"]:
    t = th["t"]
    h("THÈME %d — %s" % (t["numero"], t["nom"].upper()), 1)
    para("Source : " + t["source"], BLEU, italic=True, size=10)
    for s in t["seances"]:
        doc.add_page_break()
        h("Séance %d / %d — %s" % (s["n"], data["total"] if MODE == "PS" else data["total"], s["titre"]), 2)
        sd = s.get("sd", "")
        para("Sous-discipline : " + sd, BLEU, bold=True, size=10)
        para("Objectif : " + T(s.get("objectif", "")), None, size=11)
        para("Support : " + T(s.get("support", "")), None, size=10)
        img = s.get("image")
        if img and os.path.exists(os.path.join(ASSETS, img["file"])):
            doc.add_picture(os.path.join(ASSETS, img["file"]), width=Inches(6.0))
            para(img.get("legende", ""), BLEU, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        lc = s.get("lecon")
        if lc:
            h("Leçon : " + lc["titre"], 3, VERT)
            for b in lc.get("blocs", []):
                p = doc.add_paragraph(); r = p.add_run("• " + b["st"] + " : "); setfont(r, None, True, size=11)
                r2 = p.add_run(T(b["txt"])); setfont(r2, None, size=11)
            para("À retenir : " + T(lc.get("aretenir", "")), VERT, italic=True, size=11)
        for label, key in [("I. Révision", "revision"), ("II. Mise en situation", "miseEnSituation"), ("III. Présentation", "presentation"),
                           ("IV. Observation", "observation"), ("V. Analyse", "analyse"), ("VI. Synthèse", "synthese"),
                           ("VII. Application", "application"), ("VIII. Évaluation", "evaluation")]:
            stp = s.get(key)
            if not stp: continue
            para(label, ROUGE, bold=True, size=11)
            if stp.get("ens"): para("Enseignant : " + T(stp["ens"]), None, size=11)
            if stp.get("app"): para("Apprenants : " + T(stp["app"]), None, size=11)
            tech = stp.get("technique")
            if tech:
                if isinstance(tech, dict) and tech.get("sous"):
                    para(tech["sous"], BLEU, bold=True, size=11)
                for pt in (tech.get("points", []) if isinstance(tech, dict) else tech):
                    para("   - " + T(pt), None, size=11)
        if s.get("exercices"):
            h("Exercices notés", 3)
            for e in s["exercices"]:
                p = doc.add_paragraph(); r = p.add_run(e["t"] + " "); setfont(r, None, True, size=11)
                r2 = p.add_run(e["q"]); setfont(r2, None, size=11)
            para("CORRIGÉ", VERT, bold=True, size=11)
            for c in s.get("corrige", []):
                p = doc.add_paragraph()
                for seg in c:
                    if isinstance(seg, str):
                        r = p.add_run(seg); setfont(r, None, size=11)
                    else:
                        r = p.add_run(seg.get("t", "")); setfont(r, ROSE if seg.get("r") else None, seg.get("r", False), size=11)
    a = th.get("a")
    if a:
        doc.add_page_break()
        h("ANNEXES DU THÈME %d — %s" % (a["theme"], a["nom"].upper()), 2)
        h("A. Conjugaison détaillée", 3, ROUGE)
        for c in a["conjugaison"]:
            para(c["verbe"] + " — " + c["temps"], VERT, bold=True, size=12)
            para("Quand l'utiliser : " + c["quand"], None, size=11)
            eq = c["exemplesQuand"] if isinstance(c["exemplesQuand"], list) else [c["exemplesQuand"]]
            for e in eq: para("   « " + e + " »", BLEU, italic=True, size=11)
            para("Comment le former : " + c["comment"], None, size=11)
            for row in c["tableau"]:
                para("   • " + (row if isinstance(row, str) else " | ".join(row)), None, size=11)
            para("Pourquoi : " + c["pourquoi"], None, size=11)
            for px in c.get("pieges", []):
                if not (isinstance(px, (list, tuple)) and len(px) >= 2): continue
                p = doc.add_paragraph(); r = p.add_run("   ✗ " + px[0] + " → écris : "); setfont(r, None, size=11)
                r2 = p.add_run(px[1]); setfont(r2, ROSE, True, size=11)
        h("B. Grammaire", 3, ROUGE)
        for g in a["grammaire"]:
            p = doc.add_paragraph(); r = p.add_run(g["notion"] + " : "); setfont(r, None, True, size=11)
            r2 = p.add_run(g["regle"] + " Ex. : " + g["exemple"]); setfont(r2, None, size=11)
            if g.get("contre"):
                r3 = p.add_run(" Attention : " + g["contre"]); setfont(r3, ROSE, size=11)
        h("C. Orthographe", 3, ROUGE)
        for g in a["orthographe"]:
            ex = g["exemples"] if isinstance(g["exemples"], list) else [g["exemples"]]
            p = doc.add_paragraph(); r = p.add_run(g["notion"] + " : "); setfont(r, None, True, size=11)
            r2 = p.add_run(g["regle"] + " Ex. : " + ", ".join(ex)); setfont(r2, None, size=11)
            for px in g.get("pieges", []):
                if not (isinstance(px, (list, tuple)) and len(px) >= 2): continue
                para("   ✗ " + px[0] + " → " + px[1], ROSE, size=11)
        h("D. Vocabulaire", 3, ROUGE)
        for v in a["vocabulaire"]:
            p = doc.add_paragraph(); r = p.add_run(v["champ"] + " : "); setfont(r, VERT, True, size=11)
            r2 = p.add_run(v["mots"]); setfont(r2, BLEU, size=11)

out = os.path.join(HERE, "..", "output", "Manuel_Francais_9e_V2_%s.docx" % MODE)
doc.save(out)
print("Généré :", out, os.path.getsize(out) // 1024, "Ko")
