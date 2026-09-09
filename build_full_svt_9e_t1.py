# -*- coding: utf-8 -*-
"""
Module de construction Word (.docx) pour les Manuels Scolaires J-Learn
Conforme à 100 % au Skill J-Learn v20 et aux spécifications didactiques :
- Fiche de préparation : Titre FICHE DE PRÉPARATION, méta-table à bordures invisibles (2 zones),
  table de déroulement 6 colonnes officielle (Étapes & Durée, Enseignant, Apprenants, Technique, Support, Observation).
- Leçon : Titre LEÇON, titre de séance en ROUGE, sections en VERT, mots-clés en BLEU+GRAS, schémas FR légendés.
- Exercices : Titre EXERCICES, strictement en NOIR PUR (prêt pour impression).
- Corrigé : Titre CORRIGÉ avec barème détaillé.
- Exactement 1 seul sectPr en fin de document, 0 ns0.
"""

import os
import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Charte graphique officielle J-Learn
C_BLACK = RGBColor(0, 0, 0)
C_DARK = RGBColor(33, 37, 41)
C_RED = RGBColor(220, 53, 69)          # Titres de leçons
C_GREEN = RGBColor(25, 135, 84)        # Sous-titres et sections
C_EMERALD = RGBColor(13, 92, 58)       # En-têtes majeurs et thème J-Learn
C_BLUE = RGBColor(11, 94, 215)         # Mots-clés en gras
C_CORR = RGBColor(136, 14, 79)         # Titres corrigés
C_GRAY_TEXT = RGBColor(108, 117, 125)  # Légendes et métadonnées

HEX_HDR_BG = "D8EAD8"                  # Fond d'en-tête tableau de déroulement
HEX_SECTION_BG = "EAEAEA"              # Fond sections I, II, III
HEX_LIGHT_ROW = "F8F9FA"

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_border_none(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def add_p(doc, text="", bold=False, italic=False, size=10, color=C_DARK, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return p

def add_run(p, text, bold=False, italic=False, size=10, color=C_DARK):
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return run

def add_page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_break(docx.enum.text.WD_BREAK.PAGE)

# -----------------------------------------------------------------------------
# 1. FICHE DE PRÉPARATION : MÉTA-TABLE (BORDURES INVISIBLES, 2 ZONES)
# -----------------------------------------------------------------------------
def add_fiche_header_and_meta(doc, seance_num, titre, theme, obj, doc_ref, support, seance_str="1 / 33", duree_str="45 min"):
    # Titre officiel avant la méta-table
    add_p(doc, f"SÉANCE {seance_str.upper()}", bold=True, size=11, color=C_EMERALD, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=1)
    add_p(doc, f"Séance {seance_num} : {titre}", bold=True, size=12, color=C_RED, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2)
    add_p(doc, "FICHE DE PRÉPARATION", bold=True, size=13, color=C_DARK, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6)
    
    # Méta-table à 2 zones côte à côte avec bordures invisibles
    # Zone gauche (col 0: Intitulé, col 1: Valeur) | Zone droite (col 2: Intitulé, col 3: Valeur)
    t_meta = doc.add_table(rows=7, cols=4)
    t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_meta.autofit = False
    
    col_widths = [Cm(3.8), Cm(7.2), Cm(2.5), Cm(3.5)]
    for row in t_meta.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w
            
    meta_data = [
        [("Discipline :", True), ("Sciences de la Vie et de la Terre (Connaissances Usuelles)", False), ("Date :", True), ("_______", False)],
        [("Sous-discipline :", True), ("Sciences d'Observation & Hygiène", False), ("Classe :", True), ("9e (CE1)", False)],
        [("Thème :", True), (theme, False), ("Séance n° :", True), (seance_str, True)],
        [("Titre :", True), (f"Séance {seance_num} : {titre}", True), ("Durée :", True), (duree_str, False)],
        [("Objectif spécifique :", True), (obj, False), ("", False), ("", False)],
        [("Documentation :", True), (doc_ref, False), ("", False), ("", False)],
        [("Support et matériel :", True), (support, False), ("", False), ("", False)]
    ]
    
    for r_idx, r_data in enumerate(meta_data):
        row = t_meta.rows[r_idx]
        for c_idx, (txt, is_bold) in enumerate(r_data):
            cell = row.cells[c_idx]
            set_cell_border_none(cell)
            set_cell_margins(cell, top=20, bottom=20, left=40, right=40)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            if txt:
                col_t = C_EMERALD if is_bold and c_idx in [0, 2] else C_DARK
                add_run(p, txt, bold=is_bold, size=8.5, color=col_t)
                
    # Fusion des cellules de droite pour objectif, doc et support pour laisser de la place
    for r_idx in [4, 5, 6]:
        row = t_meta.rows[r_idx]
        row.cells[1].merge(row.cells[3])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# -----------------------------------------------------------------------------
# 2. TABLE DE DÉROULEMENT OFFICIELLE À 6 COLONNES
# -----------------------------------------------------------------------------
def add_deroulement_table_6cols(doc, duree_rev, rev_teacher, rev_students,
                               duree_lecon, mis_teacher, mis_students, mis_tech, mis_supp,
                               pres_teacher, pres_students, pres_tech, pres_supp,
                               obs_teacher, obs_students, obs_tech, obs_supp,
                               ana_q_ra_list, ana_tech, ana_supp,
                               syn_teacher, syn_students, syn_tech, syn_supp,
                               app_exercises_teacher, app_corriges_students, app_tech, app_supp,
                               duree_eval, eval_exercises_teacher, eval_corriges_students, eval_tech, eval_supp):
    
    # 13 lignes au total :
    # Row 0 : Super-header (Déroulement de la leçon)
    # Row 1 : En-têtes des 6 colonnes
    # Row 2 : Titre Section I. Révision
    # Row 3 : Contenu Révision
    # Row 4 : Titre Section II. NOUVELLE LEÇON
    # Rows 5 à 10 : 6 sous-étapes de la nouvelle leçon
    # Row 11 : Titre Section III. Évaluation
    # Row 12 : Contenu Évaluation
    
    total_rows = 13
    t_flow = doc.add_table(rows=total_rows, cols=6)
    t_flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_flow.autofit = False
    
    col_widths = [Cm(2.6), Cm(4.5), Cm(4.0), Cm(2.0), Cm(2.1), Cm(1.8)]
    for row in t_flow.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    # Row 0 : Super-header
    # Fusionner col 1 et 2 pour "Déroulement de la leçon"
    r0 = t_flow.rows[0]
    r0.cells[1].merge(r0.cells[2])
    p_der = r0.cells[1].paragraphs[0]
    p_der.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p_der, "Déroulement de la leçon", bold=True, size=8, color=C_EMERALD)
    set_cell_shading(r0.cells[1], HEX_HDR_BG)
    set_cell_margins(r0.cells[1], top=30, bottom=30, left=40, right=40)
    
    # Autres cellules du Row 0
    for idx in [0, 3, 4, 5]:
        set_cell_shading(r0.cells[idx], HEX_HDR_BG)
        set_cell_margins(r0.cells[idx], top=30, bottom=30, left=40, right=40)

    # Row 1 : Colonnes officielles
    r1 = t_flow.rows[1]
    headers_6 = [
        "Étapes et Durée",
        "Enseignant",
        "Apprenants",
        "Technique & Stratégie",
        "Support & Matériel",
        "Observation"
    ]
    for i, h in enumerate(headers_6):
        c = r1.cells[i]
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, h, bold=True, size=7.5, color=C_EMERALD)
        set_cell_shading(c, HEX_HDR_BG)
        set_cell_margins(c, top=35, bottom=35, left=40, right=40)

    # Row 2 : Section I. RÉVISION (fusion 6 colonnes)
    r2 = t_flow.rows[2]
    c2 = r2.cells[0]
    for ci in range(1, 6):
        c2.merge(r2.cells[ci])
    p2 = c2.paragraphs[0]
    add_run(p2, f"I. RÉVISION ({duree_rev})", bold=True, size=8, color=C_DARK)
    set_cell_shading(c2, HEX_SECTION_BG)
    set_cell_margins(c2, top=25, bottom=25, left=60, right=60)

    # Row 3 : Contenu Révision
    r3 = t_flow.rows[3]
    p0 = r3.cells[0].paragraphs[0]
    add_run(p0, f"Rappel des acquis\n({duree_rev})", bold=True, size=7.5)
    
    p1 = r3.cells[1].paragraphs[0]
    for q in rev_teacher:
        add_run(p1, f"• {q}\n", size=7.5)
        
    p2 = r3.cells[2].paragraphs[0]
    for a in rev_students:
        add_run(p2, f"• {a}\n", size=7.5)
        
    p3 = r3.cells[3].paragraphs[0]
    add_run(p3, "Q/R orales", size=7.5)
    
    p4 = r3.cells[4].paragraphs[0]
    add_run(p4, "Tableau", size=7.5)
    
    # r3.cells[5] reste vide pour Observation
    for c in r3.cells:
        set_cell_margins(c, top=30, bottom=30, left=40, right=40)

    # Row 4 : Section II. NOUVELLE LEÇON (fusion 6 colonnes)
    r4 = t_flow.rows[4]
    c4 = r4.cells[0]
    for ci in range(1, 6):
        c4.merge(r4.cells[ci])
    p4 = c4.paragraphs[0]
    add_run(p4, f"II. NOUVELLE LEÇON ({duree_lecon})", bold=True, size=8, color=C_DARK)
    set_cell_shading(c4, HEX_SECTION_BG)
    set_cell_margins(c4, top=25, bottom=25, left=60, right=60)

    # Rows 5 à 10 : 6 Sous-étapes sans durée individuelle
    substeps = [
        ("1. Mise en situation", mis_teacher, mis_students, mis_tech, mis_supp),
        ("2. Présentation", pres_teacher, pres_students, pres_tech, pres_supp),
        ("3. Observation", obs_teacher, obs_students, obs_tech, obs_supp),
        ("4. Analyse", ana_q_ra_list, None, ana_tech, ana_supp),
        ("5. Synthèse", syn_teacher, syn_students, syn_tech, syn_supp),
        ("6. Application", app_exercises_teacher, app_corriges_students, app_tech, app_supp),
    ]

    for idx, (s_name, t_act, st_act, tech, supp) in enumerate(substeps):
        row = t_flow.rows[5 + idx]
        p0 = row.cells[0].paragraphs[0]
        add_run(p0, s_name, bold=True, size=7.5)
        
        p1 = row.cells[1].paragraphs[0]
        p2 = row.cells[2].paragraphs[0]
        
        if idx == 3: # Analyse Q/RA
            for item in t_act:
                add_run(p1, f"Q: {item[0]}\n", size=7.5)
                add_run(p2, f"R: {item[1]}\n", size=7.5)
        elif isinstance(t_act, list):
            for item in t_act:
                add_run(p1, f"• {item}\n", size=7.5)
            if isinstance(st_act, list):
                for item in st_act:
                    add_run(p2, f"• {item}\n", size=7.5)
        else:
            add_run(p1, str(t_act), size=7.5)
            add_run(p2, str(st_act), size=7.5)
            
        p3 = row.cells[3].paragraphs[0]
        add_run(p3, str(tech), size=7.5)
        
        p4 = row.cells[4].paragraphs[0]
        add_run(p4, str(supp), size=7.5)
        
        # row.cells[5] reste vide pour Observation
        for c in row.cells:
            set_cell_margins(c, top=30, bottom=30, left=40, right=40)

    # Row 11 : Section III. ÉVALUATION (fusion 6 colonnes)
    r11 = t_flow.rows[11]
    c11 = r11.cells[0]
    for ci in range(1, 6):
        c11.merge(r11.cells[ci])
    p11 = c11.paragraphs[0]
    add_run(p11, f"III. ÉVALUATION ({duree_eval})", bold=True, size=8, color=C_DARK)
    set_cell_shading(c11, HEX_SECTION_BG)
    set_cell_margins(c11, top=25, bottom=25, left=60, right=60)

    # Row 12 : Contenu Évaluation
    r12 = t_flow.rows[12]
    p0 = r12.cells[0].paragraphs[0]
    add_run(p0, f"Contrôle individuel\n({duree_eval})", bold=True, size=7.5)
    
    p1 = r12.cells[1].paragraphs[0]
    for ex in eval_exercises_teacher:
        add_run(p1, f"• {ex}\n", size=7.5)
        
    p2 = r12.cells[2].paragraphs[0]
    for sol in eval_corriges_students:
        add_run(p2, f"• {sol}\n", size=7.5)
        
    p3 = r12.cells[3].paragraphs[0]
    add_run(p3, str(eval_tech), size=7.5)
    
    p4 = r12.cells[4].paragraphs[0]
    add_run(p4, str(eval_supp), size=7.5)
    
    # r12.cells[5] reste vide pour Observation
    for c in r12.cells:
        set_cell_margins(c, top=30, bottom=30, left=40, right=40)

    tblBorders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="0D5C3A"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="0D5C3A"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="0D5C3A"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="0D5C3A"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
        </w:tblBorders>
    ''')
    t_flow._tbl.tblPr.append(tblBorders)

# -----------------------------------------------------------------------------
# 3. LEÇON ÉLÈVE (CONFORME AU SKILL : LEÇON, TITRE EN ROUGE, SECTIONS EN VERT, MOTS-CLÉS EN BLEU+GRAS)
# -----------------------------------------------------------------------------
def add_lesson_page_skill(doc, titre_seance, sections_list, image_path=None, image_legend=""):
    add_page_break(doc)
    add_p(doc, "LEÇON", bold=True, size=13, color=C_DARK, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
    add_p(doc, titre_seance, bold=True, size=12, color=C_RED, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=8)
    
    for sec in sections_list:
        add_p(doc, sec["title"], bold=True, size=11, color=C_GREEN, space_before=4, space_after=2)
        for par in sec["paragraphs"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            for r_txt, r_bold, r_col in par:
                # Appliquer bleu si mot-clé en gras
                run_col = r_col if r_col else (C_BLUE if r_bold else C_DARK)
                add_run(p, r_txt, bold=r_bold, color=run_col, size=10)
                
    if image_path and os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(image_path, width=Inches(3.2))
        if image_legend:
            add_p(doc, image_legend, bold=True, italic=True, size=8.5, color=C_DARK, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)

# -----------------------------------------------------------------------------
# 4. EXERCICES ET CORRIGÉ (EXERCICES STRICTEMENT EN NOIR)
# -----------------------------------------------------------------------------
def add_exercises_and_answers_skill(doc, titre_seance, exercices_list, corriges_list):
    add_page_break(doc)
    
    # EXERCICES STRICTEMENT EN NOIR PUR
    add_p(doc, "EXERCICES", bold=True, size=13, color=C_BLACK, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
    add_p(doc, titre_seance, bold=True, size=11, color=C_BLACK, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6)
    
    for ex in exercices_list:
        add_p(doc, f"{ex['num']} — {ex['title']} ({ex['bareme']})", bold=True, size=10.5, color=C_BLACK, space_before=4, space_after=2)
        for line in ex['lines']:
            add_p(doc, line, size=10, color=C_BLACK, space_before=1, space_after=2)
            
    add_p(doc, "", space_after=6)
    add_p(doc, "CORRIGÉ", bold=True, size=13, color=C_CORR, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=2)
    add_p(doc, titre_seance, bold=True, size=11, color=C_RED, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6)
    
    for corr in corriges_list:
        add_p(doc, f"{corr['num']} — {corr['title']} ({corr['bareme']})", bold=True, size=10.5, color=C_GREEN, space_before=4, space_after=2)
        for p_runs in corr['answers']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            for r_txt, r_bold, r_col in p_runs:
                add_run(p, r_txt, bold=r_bold, color=r_col, size=10)

print("Official J-Learn Skill Builder V20 initialized.")
