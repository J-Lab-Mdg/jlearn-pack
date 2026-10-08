from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from collections import defaultdict
import copy

SRC='google_drive/emploi_du_temps_9eme.docx'
GRAY=RGBColor(0x99,0x99,0x99)

def set_cell(cell, main, counter=None):
    p=cell.paragraphs[0]
    for e in cell.paragraphs[1:]: e._element.getparent().remove(e._element)
    base=p.runs[0] if p.runs else p.add_run('')
    for r in p.runs[1:]: r._element.getparent().remove(r._element)
    base.text=main
    if counter:
        rc=p.add_run(' '+counter); rc.font.size=Pt(5); rc.font.color.rgb=GRAY; rc.bold=False

def no_borders(table):
    tblPr=table._tbl.tblPr
    for o in tblPr.findall(qn('w:tblBorders')): tblPr.remove(o)
    b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'none'); b.append(x)
    tblPr.append(b)
    for row in table.rows:
        for c in row.cells:
            tcPr=c._tc.get_or_add_tcPr()
            for o in tcPr.findall(qn('w:tcBorders')): tcPr.remove(o)
            tb=OxmlElement('w:tcBorders')
            for e in ('top','left','bottom','right'):
                x=OxmlElement('w:'+e); x.set(qn('w:val'),'nil'); tb.append(x)
            tcPr.append(tb)

def tnr_all(d):
    def tnr(r):
        r.font.name='Times New Roman'
        rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
        if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
        for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
    for p in d.paragraphs:
        for r in p.runs: tnr(r)
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs: tnr(r)
    for s in d.styles:
        try:
            if s.font is not None and s.font.name: s.font.name='Times New Roman'
        except: pass

def get_shd(cell):
    return cell._tc.get_or_add_tcPr().find(qn('w:shd'))

def set_shd(cell, shd_src):
    if shd_src is None: return
    tcPr=cell._tc.get_or_add_tcPr()
    old=tcPr.find(qn('w:shd'))
    if old is not None: tcPr.remove(old)
    tcPr.append(copy.deepcopy(shd_src))

def clear_shd(cell):
    tcPr=cell._tc.get_or_add_tcPr()
    old=tcPr.find(qn('w:shd'))
    if old is not None: tcPr.remove(old)

def build(classe, times, recre_rows, plan, vols, out, n_rows_needed, durs,
          recre_vol_text='1h 40 min (100 min) = 20 min × 5'):
    d=Document(SRC)
    t0=d.tables[0]
    set_cell(t0.rows[0].cells[0],'DREN : '); set_cell(t0.rows[1].cells[0],'CISCO : ')
    set_cell(t0.rows[2].cells[0],'ZAP : ');  set_cell(t0.rows[3].cells[0],'École : ')
    set_cell(t0.rows[0].cells[2],'Année scolaire : 2026-2027')
    set_cell(t0.rows[1].cells[2],'Code : '); set_cell(t0.rows[2].cells[2],f'Classe : {classe}')
    no_borders(t0)
    t=d.tables[1]
    shd = get_shd(t.rows[6].cells[1])
    shd = copy.deepcopy(shd) if shd is not None else None
    # inserer des clones de la ligne 12 APRES la ligne 12 (cellules non fusionnees)
    extra = max(0, n_rows_needed - 12)
    anchor = t.rows[12]._tr
    for _ in range(extra):
        clone = copy.deepcopy(t.rows[12]._tr)
        anchor.addnext(clone); anchor = clone
    # vider integralement les clones (y compris colonnes droites)
    for ri in range(13, 13+extra):
        for c in range(9): set_cell(t.rows[ri].cells[c],'')
    recre_label_row = 13+extra; total_row = 14+extra
    # pre-vider toutes les cellules de la grille (horaires + jours)
    for ri in range(1, max(12, n_rows_needed)+1):
        for c in range(0,6): set_cell(t.rows[ri].cells[c],'')
    daycol={'L':1,'M':2,'Me':3,'J':4,'V':5}
    total=defaultdict(int); seen=defaultdict(int); mins=defaultdict(int)
    for day in ['L','M','Me','J','V']:
        for ri in sorted(plan[day]): total[plan[day][ri]]+=1
    for ri in range(1, n_rows_needed+1):
        set_cell(t.rows[ri].cells[0], times.get(ri,''))
        if ri in recre_rows:
            for c in range(1,6):
                set_cell(t.rows[ri].cells[c],'RÉCRÉATION')
                set_shd(t.rows[ri].cells[c], shd)
        else:
            for c in range(1,6):
                if shd is not None: clear_shd(t.rows[ri].cells[c])
    for day in ['L','M','Me','J','V']:
        for ri in sorted(plan[day]):
            lab=plan[day][ri]; seen[lab]+=1
            set_cell(t.rows[ri].cells[daycol[day]], lab, f'{seen[lab]}/{total[lab]}')
            mins[lab]+=durs[ri]
    # colonne droite : disciplines sur les lignes 1..12
    for i in range(12):
        a,b = vols[i] if i < len(vols) else ('','')
        set_cell(t.rows[1+i].cells[7],a); set_cell(t.rows[1+i].cells[8],b)
    set_cell(t.rows[recre_label_row].cells[7],'Récréation')
    set_cell(t.rows[recre_label_row].cells[8],recre_vol_text)
    # ligne Total : cellule fusionnee unique a droite
    set_cell(t.rows[total_row].cells[7],'Total : 27h 30 min (1650 min)')
    for p in d.paragraphs:
        if 'RABE' in p.text or 'ANDRY' in p.text:
            for r in p.runs: r.text=''
    tnr_all(d)
    d.save(out)
    return mins

def report(classe, m):
    agg=defaultdict(int)
    for k,v in m.items(): agg[k.split(' – ')[0]]+=v
    print(classe, dict(agg), 'TOTAL', sum(agg.values()))
