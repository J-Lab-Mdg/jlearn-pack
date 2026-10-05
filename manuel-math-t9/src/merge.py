# Merge final — Manuel Mathématiques T9 (5 unités + annexes + sommaire complet)
from pathlib import Path
import re
from docx import Document
from docxcompose.composer import Composer
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

R = Path(__file__).resolve().parents[1]
W = R / 'work_merge'; W.mkdir(exist_ok=True)
(R / 'livrables').mkdir(exist_ok=True)
units = [R / f'Manuel_Mathematiques_T9_UNITE{i}.docx' for i in range(1, 6)]
ROMANS = {1: 'I', 2: 'II', 3: 'III', 4: 'IV', 5: 'V'}
NAMES = {1: 'Nombre', 2: 'Opération', 3: 'Algèbre', 4: 'Géométrie et Mesure', 5: 'Traitement de données'}


def text(el):
    return ''.join(el.itertext()).strip()


def stripped(src, n):
    d = Document(src); body = d.element.body; found = False
    for el in list(body):
        if el.tag == qn('w:sectPr'):
            continue
        if el.tag == qn('w:p') and f'UNITÉ {ROMANS[n]} —' in text(el) and 'SÉANCE' not in text(el):
            found = True
        if not found:
            body.remove(el)
    if not found:
        raise RuntimeError(f'Unité {n} introuvable')
    out = W / f'u{n}.docx'; d.save(out); return out


def font(run, size=12, bold=False, color=None):
    run.font.name = 'Times New Roman'; run.font.size = Pt(size); run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def head(d, s, size=18):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(s), size, True, (192, 0, 0))


def para(d, s, bold=False):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    font(p.add_run(s), 12, bold); return p


# ---------- ANNEXES ----------
a = Document()
head(a, 'ANNEXES — MATHÉMATIQUES T9', 22)
para(a, 'Synthèse des notions étudiées dans les cinq unités, en préparation de l’examen du BEPC.')
a.add_page_break(); head(a, 'ANNEXE 1 — GLOSSAIRE')
gloss = [
    ('Abscisse à l’origine', 'abscisse du point où la droite coupe l’axe horizontal (y = 0)'),
    ('Arrondi', 'valeur approchée obtenue en regardant le chiffre suivant la coupure : 0-4 on garde, 5-9 on monte'),
    ('Carré parfait', 'nombre entier qui est le carré d’un entier : 1, 4, 9, 16, 25…'),
    ('Classe médiane', 'première classe dont l’effectif cumulé atteint ou dépasse la moitié de l’effectif total'),
    ('Classe modale', 'classe de plus grand effectif d’une série groupée en intervalles'),
    ('Données primaires', 'données collectées soi-même ; les données reprises d’un organisme sont dites secondaires'),
    ('Droite des milieux', 'la droite joignant les milieux de deux côtés d’un triangle est parallèle au troisième côté et mesure sa moitié'),
    ('Échantillon', 'partie de la population réellement observée ; représentatif s’il reflète la population'),
    ('Effectif cumulé', 'somme de l’effectif d’une valeur et de tous les effectifs des valeurs inférieures'),
    ('Encadrement', 'coincement d’un nombre entre deux valeurs : 1,41 < √2 < 1,42'),
    ('Équation de droite', 'relation y = ax + b (ou forme générale) vérifiée par tous les points d’une droite'),
    ('Événement', 'sous-ensemble de l’univers ; impossible s’il est vide, certain s’il est l’univers entier'),
    ('Événement contraire', 'tout ce qui n’est pas l’événement A ; P(Ā) = 1 − P(A)'),
    ('Extrapoler', 'estimer une valeur au-delà des données connues en prolongeant la tendance'),
    ('Fonction affine', 'fonction de règle f(x) = ax + b ; son graphique est une droite'),
    ('Fonction inverse', 'fonction liant deux grandeurs de produit constant : xy = k'),
    ('Fréquence cumulée', 'effectif cumulé divisé par l’effectif total, souvent en pourcentage'),
    ('Identités remarquables', '(a + b)² = a² + 2ab + b² ; (a − b)² = a² − 2ab + b² ; (a + b)(a − b) = a² − b²'),
    ('Interpoler', 'estimer une valeur située entre deux données connues'),
    ('Médiane (statistique)', 'valeur qui partage la série rangée en deux moitiés de même effectif'),
    ('Mode', 'valeur du caractère ayant le plus grand effectif'),
    ('Nombre irrationnel', 'nombre qui ne peut pas s’écrire en fraction d’entiers : √2, π…'),
    ('Nombre rationnel', 'nombre qui peut s’écrire en fraction d’entiers ; développement décimal fini ou périodique'),
    ('Notation scientifique', 'écriture a × 10ⁿ avec 1 ≤ a < 10'),
    ('Ordonnée à l’origine', 'ordonnée du point où la droite coupe l’axe vertical (x = 0) : le b de y = ax + b'),
    ('Pente', 'taux de variation d’une droite : variation verticale ÷ variation horizontale'),
    ('Polygone régulier', 'polygone dont tous les côtés et tous les angles sont égaux ; angle au centre 360° ÷ n'),
    ('Polynôme', 'somme de monômes ; on l’additionne, le soustrait et le multiplie terme à terme'),
    ('Population', 'ensemble complet des individus concernés par une étude statistique'),
    ('Probabilité', 'nombre entre 0 et 1 mesurant la chance qu’un événement se produise'),
    ('Racine carrée', 'nombre positif dont le carré est le nombre donné : √a × √a = a'),
    ('Rapport de similitude', 'facteur k des triangles semblables : longueurs × k, aires × k², volumes × k³'),
    ('Taux', 'quotient de deux grandeurs d’unités différentes (Ar/kg, km/h) ; taux unitaire : dénominateur 1'),
    ('Taux de variation', 'variation de y divisée par la variation de x entre deux états ; le a de y = ax + b'),
    ('Thalès (propriété de)', 'une parallèle à un côté d’un triangle détermine des longueurs proportionnelles : AM/AB = AN/AC = MN/BC'),
    ('Triangles semblables', 'triangles de mêmes angles deux à deux ; leurs côtés homologues sont proportionnels'),
    ('Univers', 'ensemble de toutes les issues d’une expérience aléatoire, noté Ω'),
    ('Valeur initiale', 'valeur de la variable dépendante quand la variable indépendante vaut zéro'),
    ('Variable indépendante', 'variable que l’on choisit (x) ; la variable dépendante (y) en découle'),
    ('Variation directe', 'situation y = ax passant par l’origine ; partielle : y = ax + b avec b ≠ 0'),
]
for t, v in gloss:
    p = a.add_paragraph()
    font(p.add_run(t + ' : '), 12, True, (31, 78, 121)); font(p.add_run(v + '.'), 12)

a.add_page_break(); head(a, 'ANNEXE 2 — FORMULES ET PROPRIÉTÉS')
rows = [
    ('Notion', 'Formule ou propriété'),
    ('Sous-ensembles de ℝ', 'ℕ ⊂ ℤ ⊂ 𝔻 ⊂ ℚ ⊂ ℝ ; les irrationnels complètent ℚ dans ℝ'),
    ('Notation scientifique', 'a × 10ⁿ avec 1 ≤ a < 10 ; produit : multiplier les a, additionner les n'),
    ('Priorités', 'parenthèses → exposants → × ÷ → + − ; même priorité : de gauche à droite'),
    ('Produits en croix', 'a/b = c/d équivaut à a × d = b × c'),
    ('Lois des exposants', 'aᵐ × aⁿ = aᵐ⁺ⁿ ; aᵐ ÷ aⁿ = aᵐ⁻ⁿ ; (aᵐ)ⁿ = aᵐⁿ ; a⁰ = 1 ; √a = a^(1/2)'),
    ('Fonction affine', 'f(x) = ax + b ; a = taux de variation (pente), b = valeur initiale (ordonnée à l’origine)'),
    ('Pente par deux points', 'a = (y₂ − y₁) ÷ (x₂ − x₁)'),
    ('Droites particulières', 'horizontale : y = k ; verticale : x = k ; parallèles : même pente ; perpendiculaires : a × a′ = −1'),
    ('Fonction inverse', 'xy = k ; y = k/x ; hyperbole'),
    ('Identités remarquables', '(a+b)² = a²+2ab+b² ; (a−b)² = a²−2ab+b² ; (a+b)(a−b) = a²−b²'),
    ('Triangles semblables', 'angles égaux 2 à 2 ; côtés × k ; périmètres × k ; aires × k²'),
    ('Thalès', '(MN) ∥ (BC) ⟹ AM/AB = AN/AC = MN/BC ; attention AB = AM + MB'),
    ('Polygone régulier', 'angle au centre = 360° ÷ n ; angle intérieur = 180° − 360° ÷ n'),
    ('Volume de la pyramide', 'V = (aire de base × h) ÷ 3'),
    ('Volume du cône', 'V = (π × r² × h) ÷ 3 ; tronc = grand solide − petit solide'),
    ('Sphère et boule', 'aire = 4 × π × r² ; volume = (4 × π × r³) ÷ 3'),
    ('Agrandissement k', 'longueurs × k ; aires × k² ; volumes × k³'),
    ('Moyenne pondérée', 'x̄ = somme des (effectif × valeur) ÷ effectif total'),
    ('Médiane interpolée', 'Me = borne inf. + (N÷2 − cumul précédent) ÷ effectif de classe × amplitude'),
    ('Probabilité (équiprobabilité)', 'P = cas favorables ÷ cas possibles, entre 0 et 1'),
    ('Réunion d’événements', 'P(A ∪ B) = P(A) + P(B) − P(A ∩ B) ; incompatibles : P(A) + P(B)'),
    ('Événement contraire', 'P(Ā) = 1 − P(A)'),
]
t = a.add_table(rows=0, cols=2); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, row in enumerate(rows):
    c = t.add_row().cells
    for j, s in enumerate(row):
        font(c[j].paragraphs[0].add_run(s), 11, i == 0)

a.add_page_break(); head(a, 'ANNEXE 3 — RÉSUMÉ DES MÉTHODES')
methods = [
    ('Reconnaître un irrationnel', 'Chercher une écriture en fraction d’entiers ; développement décimal infini sans période → irrationnel (√2, π).'),
    ('Encadrer une racine carrée', 'Coincer entre deux carrés parfaits, puis resserrer rang par rang : unité, dixième, centième…'),
    ('Écrire en notation scientifique', 'Placer la virgule après le premier chiffre non nul, compter les rangs de déplacement : recul → exposant positif, avance → exposant négatif.'),
    ('Estimer un ordre de grandeur', 'Arrondir chaque terme à un nombre simple, calculer de tête, comparer au résultat exact pour déceler les erreurs.'),
    ('Utiliser la proportion', 'Poser les deux rapports dans le même ordre, vérifier ou résoudre par les produits en croix.'),
    ('Comparer avec un taux unitaire', 'Ramener chaque offre à l’unité (1 kg, 1 h, 1 L) en divisant, puis comparer les taux obtenus.'),
    ('Trouver l’équation d’une droite', 'Calculer la pente (graphique, formule des deux points ou donnée), puis b avec un point : b = y − ax.'),
    ('Tracer une droite', 'Partir de b sur l’axe vertical, appliquer la pente (avancer de 1, monter de a), relier à la règle ; contrôler avec un troisième point.'),
    ('Analyser deux droites', 'Comparer les pentes : égales → parallèles (confondues si même b) ; a × a′ = −1 → perpendiculaires ; sinon sécantes.'),
    ('Résoudre une équation du 1er degré', 'Développer, regrouper les x d’un côté, les nombres de l’autre, diviser par le coefficient, vérifier par substitution.'),
    ('Calculer avec Thalès', 'Vérifier les hypothèses (points alignés, parallèles), écrire les trois rapports depuis le sommet, résoudre par produits en croix — sans oublier AB = AM + MB.'),
    ('Mesurer l’inaccessible', 'Repérer deux triangles semblables (ombres, visées), apparier les côtés homologues, poser la proportion et résoudre.'),
    ('Calculer un volume composé', 'Découper le solide en briques simples (prisme, cylindre, pyramide, cône, boule), calculer chaque volume, additionner — ou soustraire les creux.'),
    ('Dresser un tableau cumulé', 'Additionner les effectifs en chemin ; contrôler que la dernière case égale l’effectif total ; diviser par N pour les fréquences cumulées.'),
    ('Interpoler la médiane', 'Calculer N ÷ 2, repérer la classe médiane dans les cumulés, avancer dans la classe au prorata des effectifs manquants.'),
    ('Calculer une probabilité', 'Vérifier l’équiprobabilité, compter les cas possibles puis favorables, former le quotient, contrôler qu’il reste entre 0 et 1.'),
]
for name, steps in methods:
    p = a.add_paragraph()
    font(p.add_run(name + ' — '), 12, True, (31, 78, 121)); font(p.add_run(steps), 12)

a.add_page_break(); head(a, 'ANNEXE 4 — AUTO-ÉVALUATION')
skills = [
    'Rationnels et irrationnels : reconnaissance, forme fractionnaire',
    'Valeurs exactes et approchées, arrondis, encadrements',
    'Racine carrée : carrés parfaits, encadrements successifs',
    'Sous-ensembles de ℝ et cartes de nombres',
    'Notation scientifique et préfixes SI',
    'Estimation et priorités des opérations',
    'Rapports, taux, proportions, pourcentages',
    'Lois des exposants et puissances de 10',
    'Variables, modes de représentation, fonctions affines',
    'Taux de variation, valeur initiale, règle y = ax + b',
    'Interpolation, extrapolation, fonction inverse',
    'Équations de droites : tracer, déterminer, comparer',
    'Polynômes et équations du premier degré',
    'Triangles semblables et propriété de Thalès',
    'Polygones réguliers : angles et constructions',
    'Volumes : pyramide, cône, boule, troncs, solides composés',
    'Rapport k : longueurs, aires (k²), volumes (k³)',
    'Mode, moyennes, médiane (discret et continu)',
    'Tableaux cumulés et interprétation critique',
    'Probabilités : vocabulaire, calculs, réunion, contraire',
]
t = a.add_table(rows=1, cols=4); t.style = 'Table Grid'
for i, s in enumerate(['Compétence', 'Acquis', 'En cours', 'À revoir']):
    font(t.rows[0].cells[i].paragraphs[0].add_run(s), 11, True)
for s in skills:
    c = t.add_row().cells
    font(c[0].paragraphs[0].add_run(s), 10)
    for i in range(1, 4):
        font(c[i].paragraphs[0].add_run('☐'), 14)

a.add_page_break(); head(a, 'ANNEXE 5 — INDEX')
para(a, ', '.join(sorted([x[0] for x in gloss], key=str.lower)) + '.')

a.add_page_break(); head(a, 'ANNEXE 6 — ÉVALUATIONS FORMAT EXAMEN')
para(a, 'Cinq sujets d’examen sur 20 points figurent à la fin des unités Nombre, Opération, Algèbre, Géométrie et Mesure, Traitement de données. Chaque sujet comporte cinq exercices avec corrigé détaillé et barème, dans l’esprit des épreuves du BEPC.')

a.add_page_break(); head(a, 'BIBLIOGRAPHIE ET WEBOGRAPHIE')
para(a, 'Ministère de l’Éducation Nationale de Madagascar, Programme d’études — Classe de T9, section Mathématiques (PE T9).')
para(a, 'Ministère de l’Éducation Nationale de Madagascar, Fascicule de ressources pédagogiques — Mathématiques.')
para(a, 'Collection J-Learn, Manuels de Mathématiques T6, T7 et T8, J-Lab Madagascar, édition 2026.')
para(a, 'Collection J-Learn, charte de conception des manuels scolaires, version 18.')

annex_marks = {
    'ANNEXES — MATHÉMATIQUES T9': 'annexes', 'ANNEXE 1 — GLOSSAIRE': 'annexe1',
    'ANNEXE 2 — FORMULES ET PROPRIÉTÉS': 'annexe2', 'ANNEXE 3 — RÉSUMÉ DES MÉTHODES': 'annexe3',
    'ANNEXE 4 — AUTO-ÉVALUATION': 'annexe4', 'ANNEXE 5 — INDEX': 'annexe5',
    'ANNEXE 6 — ÉVALUATIONS FORMAT EXAMEN': 'annexe6', 'BIBLIOGRAPHIE ET WEBOGRAPHIE': 'sources',
}
mark_id = 5000
for paragraph in a.paragraphs:
    if paragraph.text in annex_marks:
        start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(mark_id)); start.set(qn('w:name'), annex_marks[paragraph.text])
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(mark_id))
        paragraph._p.insert(0, start); paragraph._p.append(end); mark_id += 1
ann = W / 'annexes.docx'; a.save(ann)

# ---------- COMPOSITION ----------
master = Document(units[0]); composer = Composer(master)
for i in range(2, 6):
    composer.append(Document(stripped(units[i - 1], i)))
composer.append(Document(ann))

# ---------- SOMMAIRE COMPLET ----------
def hyperlink_paragraph(label, anchor, indent=False):
    p = OxmlElement('w:p'); pPr = OxmlElement('w:pPr')
    spacing = OxmlElement('w:spacing'); spacing.set(qn('w:after'), '80'); pPr.append(spacing)
    if indent:
        ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '360'); pPr.append(ind)
    p.append(pPr)
    h = OxmlElement('w:hyperlink'); h.set(qn('w:anchor'), anchor); h.set(qn('w:history'), '1')
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '1F4E79')
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.extend([color, u]); r.append(rPr)
    tx = OxmlElement('w:t'); tx.text = label; r.append(tx); h.append(r); p.append(h)
    return p


def sentence(s, roman):
    if 'RÉVISION' in s:
        return f'Révision de l’unité {roman}'
    if 'EXAMEN' in s:
        return f'Sujet d’examen — unité {roman}'
    return s.capitalize()


body_el = master.element.body
paras = list(master.paragraphs)
toc_i = next(i for i, p in enumerate(paras) if p.text == 'SOMMAIRE INTERACTIF')
insert_before = None
for par in paras[toc_i + 1:]:
    if par._p.xpath('.//w:br[@w:type="page"]'):
        insert_before = par._p; break

links = []
for unit_no in range(2, 6):
    links.append((f'Unité {ROMANS[unit_no]} — {NAMES[unit_no]}', f'u{unit_no}', False))
    source = Document(units[unit_no - 1])
    session_titles = [p.text for p in source.paragraphs if re.match(r'^SÉANCE \d+ / \d+ —', p.text)]
    for k, session_title in enumerate(session_titles, 1):
        raw = re.sub(r'^SÉANCE \d+ / \d+ — ', '', session_title)
        links.append((f'Séance {k} — {sentence(raw, ROMANS[unit_no])}', f'u{unit_no}_l{k}', True))
links.extend([
    ('Annexes', 'annexes', False), ('Annexe 1 — Glossaire', 'annexe1', True),
    ('Annexe 2 — Formules et propriétés', 'annexe2', True), ('Annexe 3 — Résumé des méthodes', 'annexe3', True),
    ('Annexe 4 — Auto-évaluation', 'annexe4', True), ('Annexe 5 — Index', 'annexe5', True),
    ('Annexe 6 — Évaluations format examen', 'annexe6', True), ('Bibliographie et webographie', 'sources', False),
])
if insert_before is not None:
    idx = body_el.index(insert_before)
    for label, anchor, ind in links:
        body_el.insert(idx, hyperlink_paragraph(label, anchor, ind)); idx += 1
else:
    raise RuntimeError('Point d’insertion du sommaire introuvable')

out = R / 'livrables' / 'Manuel_Mathematiques_T9_JLearn_V1.docx'
composer.save(out)
print('OK :', out, out.stat().st_size, 'octets')
