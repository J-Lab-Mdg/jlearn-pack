# -*- coding: utf-8 -*-
"""sortsy.py — rafitry ny boky : marika (signets), rohy (liens), lohateny (styles), fizarana.

Zavatra atao eto :
  * `marika(par, anarana)`  : signet Word (w:bookmarkStart/End) -> toe-javatra azo tsindriana ;
  * `lien(par, anarana, txt)` : rohy anatiny (Ctrl+clic) mankany amin'ny signet ;
  * `titre(doc, par, n)`    : style « Heading n » + outlineLvl -> miseho ao amin'ny « Navigation » ;
  * `fizarana(doc, ...)`    : fizarana vaovao (section) miaraka amin'ny lohateny miverimberina
    sy ny laharan-tsingana « Pajy x / y » (champs PAGE / NUMPAGES, fa tsy mila manavao).

Tsy mikasika ny votoatin'ny lesona akory : rafitra ihany.
"""
import docx                                  # noqa: F401
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

_LAHarana = [4000]


def _id():
    _LAHarana[0] += 1
    return _LAHarana[0]


def _el(nom, **props):
    e = OxmlElement(nom)
    for k, v in props.items():
        e.set(qn(k.replace("_", ":")), v)
    return e


def marika(par, anarana):
    """Signet autour du paragraphe `par` (nom = `anarana`)."""
    i = _id()
    debut = _el("w:bookmarkStart", w_id=str(i), w_name=anarana)
    fin = _el("w:bookmarkEnd", w_id=str(i))
    p = par._p
    pPr = p.find(qn("w:pPr"))
    if pPr is not None:
        pPr.addnext(debut)
    else:
        p.insert(0, debut)
    p.append(fin)
    return anarana


def lien(par, anarana, txt, couleur="1F4E79", gras=False, taille=None, italique=False):
    """Rohy anatiny (anchor) — tsy mila « Update field » : miseho dia mandeha."""
    h = OxmlElement("w:hyperlink")
    h.set(qn("w:anchor"), anarana)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if gras:
        rPr.append(OxmlElement("w:b"))
    if italique:
        rPr.append(OxmlElement("w:i"))
    rPr.append(_el("w:color", w_val=couleur))
    if taille:
        rPr.append(_el("w:sz", w_val=str(int(taille * 2))))
    rPr.append(_el("w:u", w_val="single"))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = txt
    r.append(rPr)
    r.append(t)
    h.append(r)
    par._p.append(h)
    return par


def _ordre(pPr):
    """Remets l'ordre des enfants de w:pPr selon le schéma OOXML (va du moteur build_fiche)."""
    try:
        import build_fiche as BF
        if hasattr(BF, "_ordre"):
            BF._ordre(pPr)
    except Exception:
        pass


def titre(doc, par, niveau=1):
    """Applique le style built-in « Heading n » + outlineLvl (volet Navigation de Word)."""
    nom = "Heading %d" % niveau
    try:
        doc.styles[nom]
    except KeyError:
        try:
            doc.styles.add_style(nom, WD_STYLE_TYPE.PARAGRAPH, builtin=True)
        except Exception:
            pass
    try:
        par.style = doc.styles[nom]
    except Exception:
        pass
    pPr = par._p.get_or_add_pPr()
    for e in pPr.findall(qn("w:outlineLvl")):
        pPr.remove(e)
    pPr.append(_el("w:outlineLvl", w_val=str(niveau - 1)))
    keep = OxmlElement("w:keepNext")
    if pPr.find(qn("w:keepNext")) is None:
        pPr.append(keep)
    _ordre(pPr)
    return par


def _champ(par, instr, defaut="1"):
    """w:fldSimple (PAGE, NUMPAGES) : calculé à l'ouverture, sans demande de mise à jour."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = defaut
    r.append(t)
    fld.append(r)
    par._p.append(fld)


def _sora(r, taille=9, gras=False, italique=False, loko="595959"):
    r.font.size = Pt(taille)
    r.bold = gras
    r.italic = italique
    r.font.color.rgb = RGBColor.from_string(loko)
    r.font.name = "Times New Roman"
    return r


def configurer_pied_et_tete(section, tete=None, signature=True, premier=False):
    """En-tête (nom du faritra) + pied de page (Pajy x / y) pour une section donnée."""
    if not premier:
        section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    for r in list(hp.runs):
        r._r.getparent().remove(r._r)
    if tete:
        _sora(hp.add_run(tete), taille=9, italique=True)
        pPr = hp._p.get_or_add_pPr()
        if pPr.find(qn("w:pBdr")) is None:
            bdr = OxmlElement("w:pBdr")
            btm = _el("w:bottom", w_val="single", w_sz="4", w_space="2", w_color="BFBFBF")
            bdr.append(btm)
            pPr.insert(0, bdr)
        _ordre(pPr)
    fp = section.footer.paragraphs[0]
    for r in list(fp.runs):
        r._r.getparent().remove(r._r)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if signature:
        _sora(fp.add_run("J-Learn · Malagasy T6 — "), taille=9, loko="7F7F7F")
    if signature:
        _sora(fp.add_run("Pajy "), taille=9)
    _champ(fp, " PAGE ", "1")
    _sora(fp.add_run(" / "), taille=9)
    _champ(fp, " NUMPAGES ", "1")


def fizarana(doc, tete=None, signature=True, premier=False):
    """Nouvelle section (nouvelle page) + en-tête/pied propres ; renvoie la section."""
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    configurer_pied_et_tete(sec, tete=tete, signature=signature, premier=premier)
    return sec


def table_sommaire(doc, lignes, largeurs=(560, 5300, 2200, 1800)):
    """Tabilao feno rohy : chaque ligne = (col1, col2, col3, col4, anara-signet|None)."""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    import build_fiche as BF
    t = doc.add_table(rows=1 + len(lignes), cols=4)
    BF.style_table(t, BF.STYLE_GRILLE)
    BF.largeur(t, list(largeurs), total=sum(largeurs))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(("n°", "ZANA-TARANJA sy LOHATENY", "Laharana ao amin'ny FRP",
                           "Herinandro")):
        BF.cellule(t.cell(0, j), h, bold=True)
    for i, (c1, c2, c3, c4, anc) in enumerate(lignes, start=1):
        BF.cellule(t.cell(i, 0), c1, bold=True)
        cell = t.cell(i, 1)
        par = cell.paragraphs[0]
        if anc:
            lien(par, anc, c2)
        else:
            par.add_run(c2)
        BF.cellule(t.cell(i, 2), c3)
        BF.cellule(t.cell(i, 3), c4)
    return t
